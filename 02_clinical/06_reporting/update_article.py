from pathlib import Path
import math
import pandas as pd
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(r"C:\Users\A\Documents\Codex\2026-08-31\wo-z")
OUT = ROOT / "outputs/chapter5_0831"
SRC = ROOT / "work/chapter5_0831/doc/source_0830_7.docx"
DST = ROOT / "work/chapter5_0831/doc/0831_1_静态与动态生存预测模型比较研究_第五章完整重算_0831.docx"

doc = Document(SRC)

def find_para(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def set_para(prefix, text, style=None):
    p = find_para(prefix)
    p.clear(); p.add_run(text)
    if style: p.style = style
    return p

def add_after(p, text="", style=None):
    new = OxmlElement("w:p")
    p._p.addnext(new)
    q = Paragraph(new, p._parent)
    if text: q.add_run(text)
    if style: q.style = style
    return q

def picture_paragraph_before(p):
    prev = p._p.getprevious()
    if prev is not None and prev.tag == qn("w:p"):
        return Paragraph(prev, p._parent)
    new = OxmlElement("w:p")
    p._p.addprevious(new)
    return Paragraph(new, p._parent)

def set_picture(p, path, width=6.2):
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    return p

def replace_figure(caption_prefix, image_name, new_caption=None, width=6.2):
    cap = find_para(caption_prefix)
    prev = cap._p.getprevious()
    if prev is not None:
        pic = Paragraph(prev, cap._parent)
    else:
        pic = add_after(cap, "")
    set_picture(pic, OUT / "figures" / image_name, width)
    if new_caption:
        cap.clear(); cap.add_run(new_caption); cap.style = "图片标题"
    return cap

def insert_figure_after(anchor, image_name, caption, width=6.2):
    pic = add_after(anchor)
    set_picture(pic, OUT / "figures" / image_name, width)
    cap = add_after(pic, caption, "图片标题")
    return cap

def cell_text(cell, text):
    cell.text = str(text)
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def f4(x):
    return "—" if pd.isna(x) else f"{float(x):.4f}"

def pval(x):
    x=float(x)
    return "<0.001" if x < .001 else f"{x:.3f}"

static = pd.read_csv(OUT / "tables/table_5_4A_static_metrics.csv").set_index("model")
cox = pd.read_csv(OUT / "tables/table_5_5_cox_coefficients.csv").set_index("term")
vimp = pd.read_csv(OUT / "tables/rsf_OOB_VIMP.csv").set_index("variable")
depth = pd.read_csv(OUT / "tables/rsf_minimal_depth.csv").set_index("variable")
shap = pd.read_csv(OUT / "tables/rsf_SHAP_importance.csv").set_index("variable")
dyn = pd.read_csv(OUT / "tables/table_5_4B_dynamic_metrics.csv").set_index("model")
assoc = pd.read_csv(OUT / "tables/table_5_6_JM_association.csv").set_index("trajectory")
rvimp = pd.read_csv(OUT / "tables/table_5_7_rsflc_OOB_VIMP.csv")
typ = pd.read_csv(OUT / "tables/rsf_typical_patients.csv")

# Chapter 5 data construction and locked analysis protocol.
set_para("主要结局为入院后28天死亡。院内死亡", "主要结局为入院后28天全因死亡。原始6299次住院中，337例28天死亡记录因精确死亡时刻缺失、非正或超过28天而无法用于统一生存时间分析，故从四模型共同分析数据中排除。静态任务训练集/验证集分别为4154/1685例（死亡665/282例）；第5天landmark任务训练集/验证集分别为3769/1524例（第5～28天死亡332/136例）。")
set_para("MIMIC-IV中的纵向临床数据", "MIMIC-IV纵向数据首先形成患者—日结构。静态任务的时间起点为入院，使用11个固定临床基线变量和20个纵向指标的第1天截面，共31个变量；动态任务以第5天为landmark，仅使用截至第5天可获得的信息，JM与RSFLC统一采用GCS、SOFA总分、CNS三条轨迹和25个固定协变量。")
set_para("静态任务在相同的入院/第1天信息", "静态任务在相同的31变量集合和相同验证名单上比较COX与RSF。COX采用Efron法处理并列事件时刻；RSF固定使用ntree=500、mtry=3、nodesize=10、nsplit=10。两模型均在group=1训练集拟合，在group=2独立验证集计算C-index、28天时间依赖AUC、28天IPCW Brier和IBS(0～28天)。")
set_para("训练—验证划分应按患者标识", "训练—验证划分沿用现有脚本的group变量，并在模型间保持完全一致。本次不重新执行超参数筛选，而将现有脚本所用组合视为暂定最优组合，避免把独立验证集用于调参。")
set_para("COX与RSF使用相同的入院基线变量", "COX与RSF均使用31个变量：11个固定临床基线变量，以及20个纵向指标在第1天的截面值。该定义与轨迹模型候选变量来源保持一致，但不引入第1天之后的信息。")
set_para("静态任务的主要指标为", "五项展示指标为训练集C-index、验证集C-index、第28天AUC、第28天Brier Score和IBS(0～28天)。非死亡患者在第28天视为已知无事件，28天AUC/Brier采用行政删失左极限权重。")
set_para("完成静态任务验证后", "最终RSF同时报告OOB收敛、训练OOB VIMP、验证集置换重要性、最小深度、PDP及SHAP；SHAP使用200例验证患者、50条随机排列路径的Monte Carlo近似，并保留高风险死亡与低风险存活典型患者的局部解释。")
set_para("RSF超参数在静态任务训练集内部选择", "本次未重新执行RSF网格筛选；按研究计划将现有脚本组合ntree=500、mtry=3、nodesize=10、nsplit=10视为暂定最优，并用独立验证集完成最终性能评价。")
set_para("RSF候选网格保留原设定", "表5-3A保留候选搜索空间用于后续正式筛选；本版只锁定固定组合，不把留出验证结果回填为五折筛选性能。")
set_para("表5-3A", "表5-3A 静态任务RSF暂定超参数组合")
set_para("注：搜索空间保留", "注：本次未重新筛选超参数，五折栏不报告数值；表5-4A为固定组合在独立验证集上的结果。")

# Table 5-3A and 5-4A.
t=doc.tables[8]
cell_text(t.cell(1,2),"ntree=500; mtry=3; nodesize=10; nsplit=10")
for j in range(3,6): cell_text(t.cell(1,j),"本次未重筛")
t=doc.tables[9]
for j,h in enumerate(["模型","模型设定","训练集 C-index","验证集 C-index","28天AUC","28天Brier","IBS 0–28"]): cell_text(t.cell(0,j),h)
for ri,m in [(1,"COX"),(2,"RSF")]:
    row=static.loc[m]
    setting="31变量（11个固定基线+20个第1天截面）" if m=="COX" else "ntree=500, mtry=3, nodesize=10, nsplit=10"
    vals=[m,setting,row.train_C_index,row.valid_C_index,row.AUC_28,row.Brier_28,row.IBS_0_28]
    for j,x in enumerate(vals): cell_text(t.cell(ri,j), f4(x) if j>=2 else x)
set_para("COX与RSF均使用8个基线变量", "COX与RSF均使用31个变量在同一训练集拟合并在同一验证集评价。COX验证C-index、28天AUC、Brier和IBS分别为0.7931、0.8200、0.1166和0.0943；RSF分别为0.8108、0.8375、0.1100和0.0886。RSF在区分度和预测误差上均略优于COX。")
set_para("COX验证集C-index和28天AUC", "1000次患者层面bootstrap显示，COX验证C-index的95%CI为0.7712～0.8151，RSF为0.7900～0.8314；两模型区间存在重叠，因此差异应解释为本队列内的预测增益，而非已证实的普遍优势。", "Body Text")
set_para("注：COX训练集C-index", "注：COX训练集C-index为表观值，RSF训练集C-index为OOB估计；模型排序以group=2独立验证结果为主。Brier与IBS均采用IPCW；无事件者在行政随访终点第28天仍作为已知存活对照。COX Schoenfeld残差全局检验P<0.001，提示比例风险假设整体不成立。")

# Table 5-5 COX.
label_term={"机械通气":"mechvent1","择期手术":"electivesurgery1","年龄":"age","Charlson合并症指数":"charlson_comorbidity_index","APSIII":"apsiii","SAPSII":"sapsii","OASIS":"oasis","GCS（第1天）":"gcs","血糖（第1天）":"glucose","血红蛋白（第1天）":"hemoglobin","pH（第1天，每增加0.01）":"ph"}
t=doc.tables[10]
for i in range(1,len(t.rows)):
    lab=t.cell(i,0).text.strip(); term=label_term[lab]; r=cox.loc[term]
    if term=="ph":
        hr=math.exp(float(r.beta)*.01); lo=math.exp(math.log(float(r.lower95))*.01); hi=math.exp(math.log(float(r.upper95))*.01)
    else: hr,lo,hi=float(r.HR),float(r.lower95),float(r.upper95)
    cell_text(t.cell(i,1),f"{hr:.3f}（{lo:.3f}～{hi:.3f}）"); cell_text(t.cell(i,2),pval(r.p_value))
set_para("COX应与RSF使用相同变量定义", "COX与RSF使用相同31变量。机械通气与死亡风险升高关联最强；年龄、CCI、APSIII和SAPSII亦呈正向关联。由于全局比例风险检验未通过，表5-5的HR用于描述训练集中的平均关联，不应解释为随时间恒定的因果效应。")
set_para("注：HR、95%CI及P值待", "注：HR为多变量调整结果；pH按每增加0.01单位换算。全局比例风险检验χ²=104.22，P<0.001。")

# Table 5-5A.
var_labels={"机械通气（mechvent）":"mechvent","SAPS II":"sapsii","OASIS":"oasis","APS III":"apsiii","年龄":"age","Charlson合并症指数":"charlson_comorbidity_index","24小时CNS评分":"cns_24hours","GCS":"gcs","24小时呼吸评分":"respiration_24hours","24小时SOFA评分":"sofa_24hours"}
t=doc.tables[11]
for i in range(1,len(t.rows)):
    lab=t.cell(i,0).text.strip(); nm=var_labels[lab]
    cell_text(t.cell(i,1),f4(vimp.loc[nm,"OOB_VIMP"])); cell_text(t.cell(i,2),f"{depth.loc[nm,'minimal_depth']:.3f}"); cell_text(t.cell(i,3),f4(shap.loc[nm,"mean_abs_SHAP"]))
set_para("OOB误差随树数增加", "500棵树时RSF训练OOB C-index为0.7994。OOB VIMP、最小深度和SHAP共同识别机械通气、SAPSII、OASIS、APSIII、年龄和CCI为主要预测信息；不同解释指标衡量的对象不同，排序不要求完全一致。")
set_para("最小深度结果进一步显示", "最小深度最小的变量依次包括OASIS（3.102）、SAPSII（3.246）、APSIII（3.424）、机械通气（3.446）、CCI（3.510）和GCS（3.604），提示这些变量较早进入树的分裂结构。")
set_para("PDP显示", "PDP显示机械通气对应更高的平均28天风险，SAPSII、OASIS和APSIII随数值升高呈总体上升且非线性的风险关系。PDP为模型平均响应，不代表个体因果效应。")
set_para("在验证集中选取200例患者", "验证集200例、50条随机排列路径的Monte Carlo SHAP结果中，平均绝对SHAP居前的变量为机械通气（0.0384）、SAPSII（0.0306）、OASIS（0.0227）、APSIII（0.0193）、年龄（0.0159）和CCI（0.0140）。")
set_para("局部解释选择验证集中预测风险最高", f"局部解释选择验证集中预测风险最高的死亡患者和风险最低的28天存活患者；两者预测风险分别为{typ.iloc[0].risk28:.4f}和{typ.iloc[1].risk28:.4f}。个体SHAP仅解释模型预测的组成，不等同于治疗效应。")

replace_figure("图5-2", "RSF_02_OOB_VIMP.png", "图5-2 RSF置换变量重要性（训练集OOB）")
replace_figure("图5-3", "RSF_05_PDP.png", "图5-3 RSF主要变量的28天死亡风险部分依赖")
replace_figure("图5-4", "RSF_06_SHAP_beeswarm.png", "图5-4 验证集全局SHAP分布（n=200，nsim=50）")
cap5=replace_figure("图5-5", "RSF_07_typical_patient_SHAP.png", "图5-5 高风险死亡与低风险存活典型患者的局部SHAP解释",5.8)
insert_figure_after(cap5,"RSF_08_SHAP_dependence.png","图5-6 RSF首要变量的SHAP依赖关系")

# Dynamic task.
set_para("动态任务以入住ICU后第5天为landmark", "动态任务以入院后第5天为landmark。仅纳入第5天仍存活且未删失、三条轨迹各至少有2次观测者；训练/验证分别为3769/1524例，第5～28天死亡332/136例。预测目标为在已存活至第5天条件下的第28天死亡风险。")
set_para("JM与RSFLC仅使用0～5天", "JM与RSFLC统一使用截至第5天的GCS、SOFA总分和CNS三条轨迹，以及25个固定协变量；任何第5天后的测量均不参与模型拟合或预测。RSFLC采用DynForest生存结局模式，而非28天二分类森林。")
set_para("RSFLC将在动态任务训练集内重新调参", "本次不重新执行RSFLC网格筛选，固定采用现有脚本组合ntree=200、mtry=3、nodesize=1、minsplit=2和t0=5。")
set_para("表5-3B", "表5-3B 动态任务RSFLC暂定超参数组合")
t=doc.tables[12]; cell_text(t.cell(1,2),"ntree=200; mtry=3; nodesize=1; minsplit=2; t0=5")
for j in range(3,6): cell_text(t.cell(1,j),"本次未重筛")
set_para("JM与RSFLC将在同一第5天风险队列", "JM与landmark生存RSFLC在同一第5天风险队列评价。JM验证C-index、条件AUC、Brier和IBS(5～28)分别为0.6870、0.6964、0.0889和0.0654；RSFLC分别为0.8006、0.8136、0.0739和0.0558。RSFLC在本队列内的动态预测性能优于JM，但训练C-index 0.9634提示较强拟合，外部验证仍不可缺少。")
set_para("表5-4B", "表5-4B 动态任务内部留出验证集性能")
t=doc.tables[13]; cell_text(t.cell(0,6),"IBS 5–28")
for ri,m in [(1,"JM"),(2,"RSFLC")]:
    r=dyn.loc[m]; vals=[m,r.setting,r.train_C_index,r.valid_C_index,r.AUC_28,r.Brier_28,r.IBS_5_28]
    for j,x in enumerate(vals): cell_text(t.cell(ri,j),f4(x) if j>=2 else x)

set_para("JM的纵向标记、随机效应结构", "JM采用三条链、每链1500次迭代并丢弃前500次。SOFA关联参数Rhat=1.033，接近收敛；GCS与CNS的Rhat分别为2.380和1.647，且ESS偏低，说明用户指定的短链设置不足以支持稳定参数解释。表5-6仍报告本次后验摘要，但不据此作确定性临床结论。")
t=doc.tables[14]
assoc_map={1:"value(gcs)",2:"value(sofa_24hours)",3:"value(cns_24hours)"}
for i,nm in assoc_map.items():
    r=assoc.loc[nm]; cell_text(t.cell(i,1),f"{r.alpha:.4f}"); cell_text(t.cell(i,2),f"{r.HR_per_unit:.3f}（{r.lower95:.3f}～{r.upper95:.3f}）"); cell_text(t.cell(i,3),f"{r.Rhat:.3f}")
set_para("注：后验均数、HR", "注：HR=exp(α)，95%区间为后验可信区间。通常Rhat<1.05方可认为收敛良好；本次GCS和CNS未达到该标准。")
set_para("动态任务结果只有在JM通过收敛诊断", "动态预测性能是按当前短链JM和固定超参数RSFLC得到的内部验证结果；JM关联参数未完全收敛，RSFLC训练—验证差距较大。因此表5-4B可用于本次算法实现比较，不应替代更长链复核、外部验证和临床效用评估。")

# Dynamic interpretation and table 5-7.
set_para("动态RSFLC的变量重要性、SHAP值和典型患者结果待", "RSFLC原生OOB VIMP以变量被扰动后IBS增加量衡量重要性；表5-7列出前9项。验证集置换分析使用固定种子抽取50例、每变量3次重复，作为探索性内部检查。", "Body Text")
placeholders=[p for p in doc.paragraphs if p.text.strip().startswith("动态RSFLC的变量重要性、SHAP值和典型患者结果待")]
if placeholders:
    placeholders[0].clear(); placeholders[0].add_run("固定协变量SHAP采用验证集12例、3次Monte Carlo近似；三条轨迹保持患者自身历史不变，因此该SHAP图只解释25个固定协变量，不替代轨迹层面的原生VIMP。")
if len(placeholders)>1:
    placeholders[1].clear(); placeholders[1].add_run("典型患者局部SHAP分别展示高风险死亡和低风险存活个体；依赖图展示固定协变量值与SHAP贡献的关系。所有解释均针对第5天条件风险。")
set_para("静态RSF的变量重要性、SHAP值和典型患者结果待", "RSF与RSFLC的重要性尺度、风险集和预测起点均不同。RSFLC的VIMP、置换重要性和SHAP只在动态任务内部解释，不与静态RSF的绝对数值直接比较。")
set_para("表5-7", "表5-7 Landmark生存RSFLC原生OOB变量重要性")
t=doc.tables[15]
for i in range(1,len(t.rows)):
    r=rvimp.iloc[i-1]; cell_text(t.cell(i,0),r.variable); cell_text(t.cell(i,1),f4(r.OOB_VIMP))
set_para("注：两种模型的重要性尺度不同", "注：RSFLC VIMP为扰动变量后OOB IBS的增加量，数值越大表示变量对动态预测越重要；负值表示扰动未造成误差增加，不作保护效应解释。")

anchor=find_para("动态RSFLC的变量重要性、SHAP值和典型患者结果待") if any(p.text.strip().startswith("动态RSFLC的变量重要性、SHAP值和典型患者结果待") for p in doc.paragraphs) else find_para("5.5.1")
for image,caption,width in [
    ("RSFLC_01_OOB_VIMP.png","图5-7 Landmark生存RSFLC原生OOB变量重要性",6.2),
    ("RSFLC_02_validation_permutation_VIMP.png","图5-8 RSFLC验证集置换变量重要性（n=50，每变量3次）",6.2),
    ("RSFLC_03_SHAP_beeswarm.png","图5-9 RSFLC固定协变量SHAP分布（n=12，nsim=3）",6.2),
    ("RSFLC_04_typical_patient_SHAP.png","图5-10 RSFLC典型患者局部SHAP解释",5.8),
    ("RSFLC_05_SHAP_dependence.png","图5-11 RSFLC首要固定协变量SHAP依赖关系",6.2),
    ("JM_03_MCMC_trace.png","图5-12 JM关联参数MCMC轨迹图",6.0),
    ("JM_04_MCMC_density.png","图5-13 JM关联参数后验密度图",6.0),
]:
    if (OUT/"figures"/image).exists(): anchor=insert_figure_after(anchor,image,caption,width)

# Calibration and DCA.
set_para("静态任务校准图与决策曲线待", "静态任务在同一验证集比较COX与RSF的28天校准和决策曲线。")
cap=find_para("图5-6 测试集28天校准图")
prev=picture_paragraph_before(cap); set_picture(prev,OUT/"figures"/"STATIC_09_calibration.png",6.2)
cap.clear();cap.add_run("图5-14 静态任务验证集28天校准图");cap.style="图片标题"
insert_figure_after(cap,"STATIC_10_DCA.png","图5-15 静态任务验证集28天死亡风险决策曲线")
set_para("动态任务校准图与决策曲线待", "动态任务分别展示JM与RSFLC在第5天landmark条件下的28天校准和决策曲线；鉴于JM链未完全收敛，曲线仅反映本次拟合。")
cap=find_para("图5-7 测试集28天死亡风险决策曲线")
prev=picture_paragraph_before(cap); set_picture(prev,OUT/"figures"/"RSFLC_06_calibration.png",6.2)
cap.clear();cap.add_run("图5-16 RSFLC验证集第5天landmark校准图");cap.style="图片标题"
cap=insert_figure_after(cap,"JM_01_calibration.png","图5-17 JM验证集第5天landmark校准图")
cap=insert_figure_after(cap,"RSFLC_07_DCA.png","图5-18 RSFLC动态决策曲线")
cap=insert_figure_after(cap,"JM_02_DCA.png","图5-19 JM动态决策曲线")
if (OUT/"figures"/"ALL_01_time_dependent_metrics.png").exists(): insert_figure_after(cap,"ALL_01_time_dependent_metrics.png","图5-20 四模型时间依赖AUC与Brier曲线",6.4)

set_para("本章讨论围绕两项组内比较展开", "静态任务中，RSF相较COX获得小幅区分度提升并降低Brier/IBS，解释结果一致指向机械通气、严重程度评分、年龄和合并症负担。动态任务中，landmark生存RSFLC明显优于当前短链JM，但其训练C-index很高，提示需要外部验证和进一步正则化检查。")
set_para("本段结果待按双任务协议", "两组比较不能跨任务直接排序：静态模型从入院起预测，动态模型以存活至第5天为条件。JM关联参数的未收敛和337例死亡时刻不完整记录是本版最重要的限制；后续应核验死亡时间、延长MCMC链并开展时间外或外部验证。")
set_para("主要结局统一为入院后28天死亡", "主要结局统一为入院后28天全因死亡。静态任务从入院/第1天开始预测；动态任务以第5天为条件起点，只统计第5～28天死亡。337例28天死亡记录因精确死亡时刻缺失、非正或超过28天而从四模型统一生存分析中排除。")
set_para("静态任务的最终讨论待模拟", "静态任务结果显示RSF较COX有小幅预测增益，并揭示非线性风险关系；但COX比例风险假设整体未满足，RSF解释仍限于预测关联。")
set_para("动态任务的最终讨论待JM", "动态任务结果显示landmark生存RSFLC优于当前短链JM；JM部分关联参数未收敛且RSFLC存在明显训练—验证差距，因此结论需以更长链和外部验证复核。")

# Remove the obsolete duplicated table caption retained by the source layout.
for p in list(doc.paragraphs):
    if p.text.strip() == "表5-3A 静态任务RSF超参数选择（性能待重算）":
        p._element.getparent().remove(p._element)

# Request field updates when Word opens the document.
settings=doc.settings._element
upd=settings.find(qn("w:updateFields"))
if upd is None:
    upd=OxmlElement("w:updateFields"); settings.append(upd)
upd.set(qn("w:val"),"true")

doc.save(DST)
print(DST)
