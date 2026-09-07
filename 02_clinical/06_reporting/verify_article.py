from pathlib import Path
from docx import Document
from zipfile import ZipFile

p=Path(r"work/chapter5_0831/doc/0831_1_静态与动态生存预测模型比较研究_第五章完整重算_0831.docx")
with ZipFile(p) as z:
    bad=z.testzip()
    print("ZIP_OK", bad is None, "MEDIA", len([n for n in z.namelist() if n.startswith("word/media/")]))
d=Document(p)
text="\n".join(x.text for x in d.paragraphs)
for token in ["待重算","待填写","待最终","暂不填写","性能待重算"]:
    print(token, text.count(token))
    for i,x in enumerate(d.paragraphs):
        if token in x.text: print("  HIT",i,x.text)
for prefix in ["表5-4A","表5-4B","表5-6","表5-7","图5-20"]:
    print(prefix, sum(x.text.strip().startswith(prefix) for x in d.paragraphs))
print("PARAGRAPHS",len(d.paragraphs),"TABLES",len(d.tables),"INLINE_SHAPES",len(d.inline_shapes))
for ti in [9,13,14,15]:
    print("TABLE",ti)
    for row in d.tables[ti].rows:
        print(" | ".join(c.text.replace("\n","/") for c in row.cells))
