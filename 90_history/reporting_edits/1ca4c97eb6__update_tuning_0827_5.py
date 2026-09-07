# -*- coding: utf-8 -*-
"""用更新后的五折调参结果刷新表5-3，另存 0827_5_*.docx"""
from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"F:\文章_大论文")
SRC = ROOT / r"0722\大论文版本\0827_4_基于机器学习的生存分析方法探讨_第五章超参数选择.docx"
OUT = ROOT / r"0722\大论文版本\0827_5_基于机器学习的生存分析方法探讨_第五章超参数更新.docx"
RSF_DIR = ROOT / r"0722\实例研究代码\5折RSF"
RSFLC_DIR = ROOT / r"0722\实例研究代码\5折RSFLC"

HEADER_FILL = "D9E2F3"
FONT_ASCII = "Times New Roman"
FONT_EA = "宋体"
FONT_SIZE_PT = 10.5


def set_run_font(run, bold: bool | None = None):
    run.font.name = FONT_ASCII
    run.font.size = Pt(FONT_SIZE_PT)
    if bold is not None:
        run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), FONT_ASCII)
    rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    rFonts.set(qn("w:eastAsia"), FONT_EA)
    rFonts.set(qn("w:cs"), FONT_ASCII)


def set_cell_text(cell, text: str, *, bold: bool = False, align: str = "center"):
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        set_run_font(p0.runs[0], bold=bold)
        for r in p0.runs[1:]:
            r.text = ""
    else:
        run = p0.add_run(text)
        set_run_font(run, bold=bold)
    pPr = p0._p.get_or_add_pPr()
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), align)


def set_cell_shading(cell, fill: str | None):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tcPr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def apply_header_format(table):
    for ci, cell in enumerate(table.rows[0].cells):
        set_cell_shading(cell, HEADER_FILL)
        set_cell_text(cell, cell.text.strip(), bold=True, align="center")
    for ri, row in enumerate(table.rows[1:], start=1):
        for ci, cell in enumerate(row.cells):
            set_cell_shading(cell, None)
            set_cell_text(
                cell,
                cell.text.strip(),
                bold=False,
                align="left" if ci <= 2 else "center",
            )


def replace_paragraph_text(paragraph, new_text: str):
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def read_fold_csvs(folder: Path) -> pd.DataFrame:
    dfs = []
    for i in range(1, 6):
        p = folder / f"{i}.csv"
        try:
            df = pd.read_csv(p, encoding="utf-8")
        except UnicodeDecodeError:
            df = pd.read_csv(p, encoding="gbk")
        df["fold"] = i
        dfs.append(df)
    return pd.concat(dfs, ignore_index=True)


def fmt_pm(mean: float, sd: float, digits: int = 3) -> str:
    return f"{mean:.{digits}f} ± {sd:.{digits}f}"


def summarize():
    rsf = read_fold_csvs(RSF_DIR)
    rsflc = read_fold_csvs(RSFLC_DIR)
    rsf_g = (
        rsf.groupby(["combo", "ntree", "mtry", "nodesize", "nsplit"], as_index=False)
        .agg(
            mean_c=("fold5_cindex", "mean"),
            sd_c=("fold5_cindex", "std"),
            mean_auc=("fold5_auc", "mean"),
            sd_auc=("fold5_auc", "std"),
            mean_b=("fold5_brier_ipcw", "mean"),
            sd_b=("fold5_brier_ipcw", "std"),
        )
        .sort_values(["mean_c", "mean_auc"], ascending=[False, False])
    )
    lc_g = (
        rsflc.groupby(["combo", "ntree", "mtry", "nodesize", "t0"], as_index=False)
        .agg(
            mean_c=("fold5_cindex", "mean"),
            sd_c=("fold5_cindex", "std"),
            mean_auc=("fold5_auc", "mean"),
            sd_auc=("fold5_auc", "std"),
            mean_b=("fold5_brier_ipcw", "mean"),
            sd_b=("fold5_brier_ipcw", "std"),
        )
        .sort_values(["mean_c", "mean_auc"], ascending=[False, False])
    )
    return rsf_g.iloc[0], lc_g.iloc[0], len(rsf_g), len(lc_g)


def main():
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    rsf_best, lc_best, n_rsf, n_lc = summarize()
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))

    rsf_param = (
        f"ntree={int(rsf_best['ntree'])}, mtry={int(rsf_best['mtry'])}, "
        f"nodesize={int(rsf_best['nodesize'])}, nsplit={int(rsf_best['nsplit'])}"
    )
    lc_param = (
        f"ntree={int(lc_best['ntree'])}, mtry={int(lc_best['mtry'])}, "
        f"nodesize={int(lc_best['nodesize'])}, t0={int(lc_best['t0'])}"
    )
    rows = [
        [
            "RSF",
            "ntree∈{300,500,1000}; mtry∈{3,6,9}; nodesize∈{10,20,30,40}; nsplit∈{10,25,50}",
            rsf_param,
            fmt_pm(rsf_best["mean_c"], rsf_best["sd_c"]),
            fmt_pm(rsf_best["mean_auc"], rsf_best["sd_auc"]),
            fmt_pm(rsf_best["mean_b"], rsf_best["sd_b"]),
        ],
        [
            "RSFLC",
            "ntree∈{50,100,200}; mtry∈{3,6,9,12}; nodesize∈{1,3,5}; t0=5",
            lc_param,
            fmt_pm(lc_best["mean_c"], lc_best["sd_c"]),
            fmt_pm(lc_best["mean_auc"], lc_best["sd_auc"]),
            fmt_pm(lc_best["mean_b"], lc_best["sd_b"]),
        ],
    ]

    # update table
    found = False
    for tbl in doc.tables:
        head = "|".join(c.text.strip() for c in tbl.rows[0].cells)
        if "最优超参数" in head and "五折验证" in head:
            # ensure 3 data rows? header + 2
            while len(tbl.rows) < 3:
                tbl.add_row()
            while len(tbl.rows) > 3:
                tr = tbl.rows[-1]._tr
                tr.getparent().remove(tr)
            for j, h in enumerate(
                ["模型", "搜索空间", "最优超参数", "五折验证 C-index", "五折验证 AUC", "五折验证 Brier"]
            ):
                set_cell_text(tbl.rows[0].cells[j], h, bold=True, align="center")
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    set_cell_text(
                        tbl.rows[i + 1].cells[j],
                        val,
                        bold=False,
                        align="left" if j <= 2 else "center",
                    )
            apply_header_format(tbl)
            found = True
            print("updated table:", head)
            break
    if not found:
        raise RuntimeError("未找到表5-3超参数表")

    # refresh intro2 paragraph that lists grid sizes / points to 表5-3
    intro2 = (
        f"RSF在11个基线协变量上搜索，网格为ntree∈{{300,500,1000}}、mtry∈{{3,6,9}}、"
        f"nodesize∈{{10,20,30,40}}、nsplit∈{{10,25,50}}，共{n_rsf}组。"
        f"RSFLC在20条纵向轨迹与11个固定协变量上搜索，网格为ntree∈{{50,100,200}}、mtry∈{{3,6,9,12}}、"
        f"nodesize∈{{1,3,5}}，预测时纵向信息截断时点固定为t0=5天，共{n_lc}组。"
        f"最优组合及其五折验证性能见表5-3。后续全训练集重拟合与独立验证集评价均采用该最优超参数。"
    )
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("RSF在11个基线协变量上搜索"):
            replace_paragraph_text(p, intro2)
            print("updated intro2")
            break

    # replace any leftover old best-param strings in body
    old_new = [
        (
            "ntree=300, mtry=3, nodesize=40, nsplit=25",
            rsf_param,
        ),
        (
            "ntree=100, mtry=3, nodesize=1, t0=5",
            lc_param,
        ),
    ]
    for p in doc.paragraphs:
        t = p.text
        nt = t
        for a, b in old_new:
            nt = nt.replace(a, b)
        if nt != t:
            replace_paragraph_text(p, nt)
            print("replaced leftover param text")

    doc.save(str(OUT))
    print("saved:", OUT)
    print("RSF:", rsf_param, rows[0][3], rows[0][4], rows[0][5])
    print("RSFLC:", lc_param, rows[1][3], rows[1][4], rows[1][5])

    doc2 = Document(str(OUT))
    for tbl in doc2.tables:
        head = "|".join(c.text.strip() for c in tbl.rows[0].cells)
        if "最优超参数" in head:
            for r in tbl.rows:
                print(" | ".join(c.text.strip()[:55] for c in r.cells))


if __name__ == "__main__":
    main()
