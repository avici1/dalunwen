# -*- coding: utf-8 -*-
"""从 0827_5 重做：删除正文旧5.3.3+5.4，新建衔接的5.3.3（验证集终评），输出0827_6。"""
from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = Path(r"F:\文章_大论文\0722\大论文版本\0827_5_基于机器学习的生存分析方法探讨_第五章超参数更新.docx")
OUT = Path(r"F:\文章_大论文\0722\大论文版本\0827_6_基于机器学习的生存分析方法探讨_第五章删除训练内五折主比较.docx")


def p_text(el) -> str:
    return "".join((t.text or "") for t in el.iter(qn("w:t"))).strip()


def set_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def clone_p_after(anchor_el, style_src_p, text: str):
    new_p = copy.deepcopy(style_src_p._p)
    texts = list(new_p.iter(qn("w:t")))
    if texts:
        texts[0].text = text
        for t in texts[1:]:
            t.text = ""
    else:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        new_p.append(r)
    anchor_el.addnext(new_p)
    return new_p


def body_children(body):
    return list(body.iterchildren())


def find_body_after(body, pred, after_el=None):
    kids = body_children(body)
    start = kids.index(after_el) + 1 if after_el is not None else 0
    for el in kids[start:]:
        if el.tag == qn("w:p") and pred(p_text(el)):
            return el
    return None


def delete_between(body, start_el, end_el_exclusive):
    kids = body_children(body)
    i0 = kids.index(start_el)
    i1 = kids.index(end_el_exclusive)
    assert i1 > i0
    for el in kids[i0:i1]:
        body.remove(el)
    print(f"deleted [{i0}:{i1}] n={i1-i0}")


def one_pass_replace(text: str, mapping: dict[str, str]) -> str:
    """按长键优先一次替换，避免连锁。"""
    if not any(k in text for k in mapping):
        return text
    # placeholder technique
    items = sorted(mapping.items(), key=lambda x: len(x[0]), reverse=True)
    tmp = text
    holders = {}
    for i, (k, v) in enumerate(items):
        if k in tmp:
            h = f"@@HOLD{i}@@"
            holders[h] = v
            tmp = tmp.replace(k, h)
    for h, v in holders.items():
        tmp = tmp.replace(h, v)
    return tmp


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    body = doc.element.body

    # 以正文「5.3.2 超参数选择结果」为锚（跳过目录：取最后一个匹配）
    anchors_532 = [el for el in body_children(body) if el.tag == qn("w:p") and p_text(el).startswith("5.3.2 超参数选择结果")]
    if not anchors_532:
        raise RuntimeError("未找到5.3.2")
    p_532 = anchors_532[-1]

    p_note = find_body_after(
        body,
        lambda t: t.startswith("注：选参准则为五折验证集平均C-index"),
        after_el=p_532,
    )
    p_old_533 = find_body_after(
        body,
        lambda t: t.startswith("5.3.3 训练集内五折交叉验证结果"),
        after_el=p_532,
    )
    p_55 = find_body_after(
        body,
        lambda t: t.startswith("5.5 模型解释"),
        after_el=p_532,
    )
    p_tab57 = find_body_after(
        body,
        lambda t: t.startswith("表5-7 独立验证集28天预测性能"),
        after_el=p_532,
    )
    if not all([p_note, p_old_533, p_55, p_tab57]):
        raise RuntimeError(f"定位失败 note={p_note is not None} old533={p_old_533 is not None} 55={p_55 is not None} tab={p_tab57 is not None}")

    # 收集表5-7题注+表+注
    kids = body_children(body)
    i_cap = kids.index(p_tab57)
    move_els = [p_tab57]
    j = i_cap + 1
    while j < len(kids) and kids[j].tag == qn("w:p") and not p_text(kids[j]):
        move_els.append(kids[j])
        j += 1
    if j >= len(kids) or kids[j].tag != qn("w:tbl"):
        raise RuntimeError("表5-7后未找到表格")
    move_els.append(kids[j])
    j += 1
    while j < len(kids) and kids[j].tag == qn("w:p") and not p_text(kids[j]):
        move_els.append(kids[j])
        j += 1
    if j < len(kids) and kids[j].tag == qn("w:p") and p_text(kids[j]).startswith("注："):
        move_els.append(kids[j])

    # 删除旧5.3.3至5.5前（含5.4整节）
    delete_between(body, p_old_533, p_55)

    # 样式源
    style_h = style_b = None
    for p in doc.paragraphs:
        if p._p is p_532:
            style_h = p
        if p.text.startswith("按照第5.3.1节的两阶段框架"):
            style_b = p
    if style_h is None or style_b is None:
        raise RuntimeError("样式源缺失")

    bridge1 = (
        "采用第5.3.2节选定的最优超参数（COX与JM按既定模型设定、无需网格调参），"
        "在全部训练集上重新拟合各模型，并在独立验证集上计算C-index、28天AUC与Brier Score等指标，"
        "作为本框架下的模型预测性能终评。该步骤不再重复使用训练集内五折交叉验证进行模型间主比较，"
        "以避免与超参数选择过程角色重叠，并与第5.3.1节“独立验证集评价”的设定保持一致。"
    )
    bridge2 = (
        "表5-4汇总独立验证集上的主要预测性能，供后续解释性分析、校准、决策曲线与风险分层共同对照。"
        "需要说明的是：不同模型可用信息与风险表达可能不完全相同（例如RSF与RSFLC的Brier计算口径），"
        "比较时以区分度指标为主，校准指标宜结合表注审慎解读。"
    )

    el = p_note
    el = clone_p_after(el, style_h, "5.3.3 最优超参数下的模型性能比较")
    el = clone_p_after(el, style_b, bridge1)
    el = clone_p_after(el, style_b, bridge2)

    for e in move_els:
        parent = e.getparent()
        if parent is not None:
            parent.remove(e)
    for e in move_els:
        el.addnext(e)
        el = e
    print("moved performance table")

    # 目录：定位靠前的条目
    toc_removed = 0
    for p in list(doc.paragraphs[:90]):
        t = p.text.strip()
        if t == "5.3.3 训练集内五折交叉验证结果":
            set_text(p, "5.3.3 最优超参数下的模型性能比较")
        elif t.startswith("5.4 模型性能的统计比较"):
            p._p.getparent().remove(p._p)
            toc_removed += 1
        elif t.startswith("5.5 模型解释"):
            set_text(p, "5.4 模型解释：变量重要性与SHAP分析\t")
        elif t.startswith("5.5.1"):
            set_text(p, t.replace("5.5.1", "5.4.1", 1))
        elif t.startswith("5.5.2"):
            set_text(p, t.replace("5.5.2", "5.4.2", 1))
        elif t.startswith("5.5.3"):
            set_text(p, t.replace("5.5.3", "5.4.3", 1))
        elif t.startswith("5.6 校准"):
            set_text(p, t.replace("5.6", "5.5", 1))
        elif t.startswith("5.7 真实数据"):
            set_text(p, t.replace("5.7", "5.6", 1))
        elif t.startswith("5.8 本章小结"):
            set_text(p, t.replace("5.8", "5.7", 1))
    print("toc_removed", toc_removed)

    # 正文标题与引用一次性映射
    mapping = {
        "5.8 本章小结": "5.7 本章小结",
        "5.7 真实数据分析结果讨论": "5.6 真实数据分析结果讨论",
        "5.6.3 ": "5.5.3 ",
        "5.6.2 ": "5.5.2 ",
        "5.6.1 ": "5.5.1 ",
        "5.6 校准、决策曲线与风险分层": "5.5 校准、决策曲线与风险分层",
        "5.5.3 ": "5.4.3 ",
        "5.5.2 ": "5.4.2 ",
        "5.5.1 ": "5.4.1 ",
        "5.5 模型解释：变量重要性与SHAP分析": "5.4 模型解释：变量重要性与SHAP分析",
        "第5.5节": "第5.4节",
        "第5.6节": "第5.5节",
        "表5-8": "表5-5",
        "表5-7": "表5-4",
        "图5-10": "图5-8",
        "图5-9": "图5-7",
        "图5-8": "图5-6",
        "图5-7": "图5-5",
        "图5-6": "图5-4",
        "图5-5": "图5-3",
        "图5-4": "图5-2",
    }

    for p in doc.paragraphs:
        t = p.text
        # 跳过已写好的新5.3.3标题
        if t.strip() == "5.3.3 最优超参数下的模型性能比较":
            continue
        nt = one_pass_replace(t, mapping)
        if nt != t:
            set_text(p, nt)

    # 题注精修
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("表5-4 独立验证集28天预测性能"):
            set_text(p, "表5-4 独立验证集28天预测性能（70%/30%划分；最优超参数下终评）")
        if t.startswith("表5-5 验证集置换变量重要性"):
            set_text(p, "表5-5 验证集置换变量重要性对照（RSFLC含GCS）")

    # 解释节导语
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("为提高机器学习模型的可解释性"):
            set_text(
                p,
                "为提高机器学习模型的可解释性，并识别卒中28天死亡预测中的主要因素，"
                "本文在独立验证集（约30%）上，基于第5.3.3节重拟合的最优RSF与RSFLC模型，"
                "同时报告置换变量重要性（permutation VIMP）与SHAP（SHapley Additive Explanations）。"
                "VIMP通过随机打乱某一变量后观察预测性能下降来衡量该变量的全局贡献；"
                "SHAP则把单次预测分解为各特征的加性贡献，正值表示推动28天死亡风险升高，负值表示推动风险降低。"
                "RSF的SHAP基于其纳入的协变量计算；RSFLC的SHAP仅扰动固定协变量，纵向轨迹保持原值，"
                "因此SHAP图不直接给出纵向轨迹变量的局部贡献。",
            )
        if t.startswith("表5-4给出两模型在独立验证集上的28天预测性能，作为本框架下的终评结果"):
            # 已前移，若解释节仍残留则删
            p._p.getparent().remove(p._p)
            print("removed leftover perf intro")

    # 框架段
    for p in doc.paragraphs:
        t = p.text
        if "表5-3给出RSF与RSFLC的超参数选择结果" in t:
            set_text(
                p,
                "在完成验证集预测后，进一步对机器学习模型（RSF、RSFLC）开展置换变量重要性（VIMP）与SHAP分析，"
                "以刻画在既定模型与评价协议下各协变量对28天死亡风险的相对贡献。"
                "该类解释性分析建立在上述最优模型及其验证集预测之上，与性能终评共用同一划分，而不另设并行评价协议。"
                "需要说明的是：表5-3给出RSF与RSFLC的超参数选择结果；"
                "表5-4给出独立验证集上的预测性能终评，并作为第5.4–5.5节解释与临床效用分析的共同对照。"
                "不以训练集内交叉验证的选参折性能替代验证集终评数字。",
            )
        if "第一阶段仅在训练集内部实施五折交叉验证，服务于两项目的" in t:
            set_text(
                p,
                t.replace(
                    "服务于两项目的：一是在候选超参数网格上，以验证折的平均C-index（必要时辅以AUC、Brier Score）筛选最优超参数组合；"
                    "二是观察各折性能波动，初步评估模型在不同训练子样本下的稳定性。"
                    "交叉验证过程不使用独立验证集，以保证终评样本独立。",
                    "主要用于在候选超参数网格上，以验证折的平均C-index（必要时辅以AUC、Brier Score）筛选最优超参数组合；"
                    "折间波动仅作为选参过程的附带信息，不作为模型间性能排序的主依据。"
                    "交叉验证过程不使用独立验证集，以保证终评样本独立。",
                ),
            )

    # 小结/讨论
    for p in list(doc.paragraphs):
        t = p.text
        if t.startswith("本章以MIMIC-IV数据库中6299次"):
            set_text(
                p,
                "本章以MIMIC-IV数据库中6299次脑卒中住院记录（对应6092名患者；含缺血性与出血性等亚型）的真实临床数据为基础，"
                "采用“训练集内五折交叉验证调参—全训练集重拟合—独立验证集评价”的框架，"
                "结合验证集上的VIMP、SHAP、校准、决策曲线和风险分层，"
                "对四种生存分析模型（COX、JM、RSF、RSFLC）的预测性能与可解释性做评价。",
            )
        if t.startswith("RSF模型展现出最高的区分度（AUC=0.7977"):
            set_text(
                p,
                "在独立验证集终评下，机器学习模型整体具有更好的区分度表现："
                "RSF的验证集C-index为0.7885、28天AUC为0.8193；"
                "RSFLC的验证集C-index为0.7736、AUC为0.8021。"
                "该结果与第5.3.2节选定的最优超参数相对应，作为本章模型比较的主要依据。"
                "有关变量贡献与临床效用，见第5.4节与第5.5节。",
            )
        if t.startswith("Friedman检验表明"):
            p._p.getparent().remove(p._p)
        if "第六，从统计检验的结果可以得出结论，交叉验证推断的局限性" in t:
            set_text(
                p,
                "第六，本章主比较建立在预先划分的独立验证集之上；"
                "训练集内五折交叉验证仅用于超参数选择，而不再作为模型间性能排序的主依据。"
                "若未来需要为验证集指标提供区间估计，可采用Bootstrap等重抽样方法，"
                "在不重复使用调参样本的前提下补充统计推断。",
            )
        if "表5-3/表5-4仅反映训练阶段选参与交叉验证稳定性" in t:
            set_text(
                p,
                t.replace(
                    "与表5-4终评及第5.4节解释性分析共用同一划分；表5-3/表5-4仅反映训练阶段选参与交叉验证稳定性，不作为本节绝对概率与净获益的引用依据。",
                    "与表5-4终评及第5.4节解释性分析共用同一划分。",
                ),
            )

    # 章首流程若仍写嵌套/五折主比较
    for p in doc.paragraphs:
        t = p.text
        if "分析流程分为五个部分" in t and "嵌套五折" not in t:
            if "训练集内五折交叉验证中完成超参数选择并评估稳定性，再以最优模型于独立验证集评价预测性能" in t:
                set_text(
                    p,
                    t.replace(
                        "在训练集内五折交叉验证中完成超参数选择并评估稳定性，再以最优模型于独立验证集评价预测性能",
                        "在训练集内五折交叉验证中完成超参数选择，再以最优超参数于全训练集重拟合并在独立验证集评价预测性能",
                    ),
                )

    doc.save(str(OUT))
    print("saved", OUT)

    doc2 = Document(str(OUT))
    print("==== structure check ====")
    for i, p in enumerate(doc2.paragraphs):
        t = p.text.strip()
        if i < 80 and t.startswith(("5.3", "5.4", "5.5", "5.6", "5.7")):
            print("TOC", i, t[:70])
    bad = []
    for i, p in enumerate(doc2.paragraphs):
        t = p.text.strip()
        if i < 200:
            continue
        if t.startswith("5.3.3 训练集内") or t.startswith("5.4 模型性能的统计比较") or "Friedman检验结果如表" in t:
            bad.append((i, t[:80]))
        if t.startswith("5.3") or t.startswith("5.4") or t.startswith("5.5") or t.startswith("表5-3") or t.startswith("表5-4") or t.startswith("采用第5.3.2"):
            if 250 <= i <= 330:
                print(i, t[:120])
    print("bad leftovers", bad)

    # figure caption uniqueness
    figs = {}
    for p in doc2.paragraphs:
        m = re.match(r"(图5-\d+)", p.text.strip())
        if m:
            figs.setdefault(m.group(1), 0)
            figs[m.group(1)] += 1
    print("fig caption counts", figs)


if __name__ == "__main__":
    main()
