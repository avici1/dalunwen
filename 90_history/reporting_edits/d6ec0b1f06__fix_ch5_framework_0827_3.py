# -*- coding: utf-8 -*-
"""补修 0827_3：正文标题、5.6导语，以及摘要/方法中残留的嵌套CV表述。"""
from pathlib import Path

from docx import Document

OUT = Path(r"F:\文章_大论文\0722\大论文版本") / (
    "0827_3_基于机器学习的生存分析方法探讨_第五章分析框架修订.docx"
)


def set_text(p, new: str) -> None:
    if p.runs:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new)


def main() -> None:
    doc = Document(str(OUT))

    set_text(doc.paragraphs[260], "5.3.1 分析框架")
    set_text(doc.paragraphs[266], "5.3.2 训练集内五折交叉验证结果")
    set_text(
        doc.paragraphs[311],
        "区分度只说明模型能否把先发生事件的患者排在前面。"
        "临床使用还取决于预测概率是否校准、在不同决策阈值下是否带来净获益，以及风险分层能否分开生存曲线。"
        "以下分析均在独立验证集、28天结局上进行，与表5-6终评及第5.5节解释性分析共用同一划分；"
        "表5-3仅反映训练集内交叉验证的稳定性，不作为本节绝对概率与净获益的引用依据。",
    )

    fixes = {
        110: (
            "模型评估采用“训练集内交叉验证调参—全训练集重拟合—独立验证集评价”的两阶段策略："
            "在训练集内以五折交叉验证结合网格搜索优化超参数并观察折间稳定性，"
            "再用最优超参数在全训练集重拟合，于独立验证集报告C指数、AUC与Brier Score；"
            "统计推断对训练集内各折性能采用Friedman检验进行多模型比较，"
            "Wilcoxon符号秩检验进行配对比较。"
            "此外，在独立验证集预测基础上采用SHAP（SHapley Additive Explanations）与置换变量重要性（VIMP）"
            "对机器学习模型进行解释，并报告校准、决策曲线与风险分层结果。"
        ),
        174: (
            "实例数据分析采用预先划分的约70%/30%训练—验证集。"
            "在训练集内实施五折交叉验证，结合超参数网格搜索，以验证折平均C-index筛选最优配置，并观察折间稳定性；"
            "再用最优超参数在全训练集重拟合，于独立验证集评估最终模型性能（C-index、AUC、Brier Score）。"
            "对RSF与RSFLC进一步在验证集上开展VIMP与SHAP解释分析。"
        ),
        181: (
            "本文整体设计可以概括为“模拟+实证”双轨验证体系。"
            "模拟研究通过控制六个核心设计维度，对4种模型在60种不同的场景下进行预测性能的系统评价，"
            "发现模型在不同数据条件下具有怎样的性能特点以及适用范围；"
            "实例研究以MIMIC-IV真实临床数据为基础，用训练集内交叉验证调参、独立验证集终评以及SHAP/VIMP解释的方法，"
            "来检验模拟研究结果能否推广到实际的应用场景中。"
            "两种研究的评估指标保持一致（C-index、AUC、Brier Score）。"
        ),
        240: (
            "本章利用MIMIC-IV数据库的脑卒中住院队列（含缺血性与出血性等亚型），对四种生存模型做真实的分析。"
            "分析流程分为五个部分，即（1）研究队列的详细描述和临床特征总结；（2）多源纵向数据预处理和融合；"
            "（3）在训练集内五折交叉验证中完成超参数选择并评估稳定性，再以最优模型于独立验证集评价预测性能；"
            "（4）在同一验证集上，对最优RSF与RSFLC模型进行置换变量重要性（VIMP）与SHAP解释；"
            "（5）基于28天死亡风险给出校准图、决策曲线分析与Kaplan–Meier风险分层。"
            "四种模型为COX、JM、RSF与RSFLC。"
        ),
    }
    for idx, new in fixes.items():
        set_text(doc.paragraphs[idx], new)
        print("updated", idx)

    # 摘要：若仍含嵌套表述则做针对性替换
    t5 = doc.paragraphs[5].text
    if "嵌套" in t5:
        nt = t5
        replacements = [
            ("采用嵌套五折交叉验证", "采用训练集内五折交叉验证调参与独立验证集评价"),
            ("用嵌套五折交叉验证", "用训练集内五折交叉验证调参与独立验证集评价"),
            ("嵌套五折交叉验证", "训练集内五折交叉验证调参与独立验证集评价"),
            ("嵌套交叉验证", "训练集内交叉验证调参与独立验证集评价"),
        ]
        for a, b in replacements:
            nt = nt.replace(a, b)
        set_text(doc.paragraphs[5], nt)
        print("updated abstract 5; still nested?", "嵌套" in nt)
        if "嵌套" in nt:
            i = nt.find("嵌套")
            print(nt[max(0, i - 50) : i + 80])

    doc.save(str(OUT))
    print("saved", OUT)

    doc2 = Document(str(OUT))
    print("==== key paras ====")
    for i in [5, 64, 65, 110, 174, 181, 240, 260, 261, 265, 266, 269, 291, 311, 330, 333]:
        print(i, doc2.paragraphs[i].text.strip()[:170])
        print("---")
    print("==== leftover 嵌套 ====")
    for i, p in enumerate(doc2.paragraphs):
        if "嵌套" in p.text:
            print(i, p.text.strip()[:140])
    print("==== leftover 不替代 ====")
    for i, p in enumerate(doc2.paragraphs):
        if "不替代表5-3" in p.text:
            print(i, p.text.strip()[:140])


if __name__ == "__main__":
    main()
