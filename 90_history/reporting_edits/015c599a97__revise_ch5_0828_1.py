# -*- coding: utf-8 -*-
"""方案B：四模型终评 + COX/JM 解释小节 + 删除5.5.3 KM；缺失/异常数字标红。"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table

ROOT = Path(r"F:\文章_大论文")
SRC = ROOT / r"0722\大论文版本\0827_6_基于机器学习的生存分析方法探讨_第五章删除训练内五折主比较.docx"
OUT = ROOT / r"0722\大论文版本\0828_1_基于机器学习的生存分析方法探讨_第五章四模型终评与KM删除.docx"
EXEC = ROOT / r"0722\实例研究代码\执行"

HEADER_FILL = "D9E2F3"
FONT_ASCII = "Times New Roman"
FONT_EA = "宋体"
FONT_SIZE_PT = 10.5
RED_VAL = "FF0000"


def p_text(el) -> str:
    if el.tag != qn("w:p"):
        return ""
    return "".join((t.text or "") for t in el.iter(qn("w:t"))).strip()


def body_children(body):
    return list(body.iterchildren())


def find_body(body, pred, after_el=None, last=False):
    kids = body_children(body)
    start = kids.index(after_el) + 1 if after_el is not None else 0
    found = None
    for el in kids[start:]:
        if el.tag == qn("w:p") and pred(p_text(el)):
            if not last:
                return el
            found = el
    return found


def set_run_font(run, bold: bool | None = None, size_pt: float = FONT_SIZE_PT, red: bool = False):
    run.font.name = FONT_ASCII
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), FONT_ASCII)
    rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    rFonts.set(qn("w:eastAsia"), FONT_EA)
    rFonts.set(qn("w:cs"), FONT_ASCII)
    for old in list(rPr.findall(qn("w:color"))):
        rPr.remove(old)
    if red:
        color = OxmlElement("w:color")
        color.set(qn("w:val"), RED_VAL)
        rPr.append(color)


def apply_segments(paragraph, segments, size_pt: float | None = None):
    """segments: list[(text, red)]. 保留原段 rPr 字号/字体，仅改颜色。"""
    p_el = paragraph._p
    rPr_template = None
    for r in p_el.findall(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            rPr_template = copy.deepcopy(rpr)
            break
    for child in list(p_el):
        if child.tag == qn("w:r"):
            p_el.remove(child)
    for text, red in segments:
        r = OxmlElement("w:r")
        if rPr_template is not None:
            rPr = copy.deepcopy(rPr_template)
        else:
            rPr = OxmlElement("w:rPr")
        for old in list(rPr.findall(qn("w:color"))):
            rPr.remove(old)
        if size_pt is not None:
            for tag in ("sz", "szCs"):
                el = rPr.find(qn(f"w:{tag}"))
                if el is None:
                    el = OxmlElement(f"w:{tag}")
                    rPr.append(el)
                el.set(qn("w:val"), str(int(size_pt * 2)))
        if red:
            color = OxmlElement("w:color")
            color.set(qn("w:val"), RED_VAL)
            rPr.append(color)
        r.append(rPr)
        t = OxmlElement("w:t")
        if text[:1].isspace() or text[-1:].isspace():
            t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p_el.append(r)


def set_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def clone_p_after(anchor_el, style_src_p, text: str | None = None, segments=None):
    new_p = copy.deepcopy(style_src_p._p)
    texts = list(new_p.iter(qn("w:t")))
    if texts:
        texts[0].text = text or (segments[0][0] if segments else "")
        for t in texts[1:]:
            t.text = ""
    else:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text or ""
        r.append(t)
        new_p.append(r)
    anchor_el.addnext(new_p)
    if segments:
        from docx.text.paragraph import Paragraph

        apply_segments(Paragraph(new_p, style_src_p._parent), segments)
    return new_p


def set_cell_shading(cell, fill: str | None):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tcPr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_v_align(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), "center")


def set_para_align(p, align: str):
    pPr = p._p.get_or_add_pPr()
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), align)


def set_cell_runs(cell, segments, *, bold: bool = False, align: str = "center", size_pt: float = FONT_SIZE_PT):
    """segments: list[(text, red)] or a plain str."""
    if isinstance(segments, str):
        segments = [(segments, False)]
    p0 = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for text, red in segments:
        run = p0.add_run(text)
        set_run_font(run, bold=bold, size_pt=size_pt, red=red)
    set_para_align(p0, align)
    set_cell_v_align(cell)


def set_tbl_border_el(borders, tag: str):
    el = borders.find(qn(f"w:{tag}"))
    if el is None:
        el = OxmlElement(f"w:{tag}")
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:color"), "000000")
    el.set(qn("w:sz"), "4")
    el.set(qn("w:space"), "0")


def apply_table_shell(table, n_cols: int):
    try:
        table.style = "Normal Table"
    except Exception:
        pass
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        set_tbl_border_el(borders, edge)

    def upsert(tag, attrs):
        el = tblPr.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tblPr.append(el)
        for k, v in attrs.items():
            el.set(qn(f"w:{k}"), v)

    upsert("tblW", {"w": "5000", "type": "pct"})
    upsert("jc", {"val": "center"})
    upsert("tblLayout", {"type": "autofit"})

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for g in list(grid.findall(qn("w:gridCol"))):
            grid.remove(g)
        width = str(int(9000 / max(n_cols, 1)))
        for _ in range(n_cols):
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), width)
            grid.append(gc)


def rebuild_table(table, headers, rows):
    """rows[i][j] = str | list[(text, red)]"""
    n_cols = len(headers)
    n_body = len(rows)
    tbl = table._tbl
    for tr in list(tbl.findall(qn("w:tr"))):
        tbl.remove(tr)

    def add_row():
        tr = OxmlElement("w:tr")
        for _ in range(n_cols):
            tc = OxmlElement("w:tc")
            tcPr = OxmlElement("w:tcPr")
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), "2000")
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)
            tc.append(tcPr)
            tc.append(OxmlElement("w:p"))
            tr.append(tc)
        tbl.append(tr)

    add_row()
    for _ in range(n_body):
        add_row()
    apply_table_shell(table, n_cols)
    wrapped = Table(tbl, table._parent)
    for j, h in enumerate(headers):
        set_cell_shading(wrapped.rows[0].cells[j], HEADER_FILL)
        set_cell_runs(wrapped.rows[0].cells[j], h, bold=True, align="center")
    for i, vals in enumerate(rows):
        for j, val in enumerate(vals):
            set_cell_shading(wrapped.rows[i + 1].cells[j], None)
            set_cell_runs(
                wrapped.rows[i + 1].cells[j],
                val,
                bold=False,
                align="left" if j == 0 else "center",
            )
    return wrapped


def insert_table_after(anchor_el, template_tbl_el, parent, headers, rows):
    new_el = copy.deepcopy(template_tbl_el)
    anchor_el.addnext(new_el)
    table = Table(new_el, parent)
    rebuild_table(table, headers, rows)
    return new_el


def fmt_p(p: float) -> str:
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def fmt_hr(lo, hi, point, nd=2) -> str:
    return f"{point:.{nd}f} ({lo:.{nd}f}–{hi:.{nd}f})"


def load_results():
    cox_perf = pd.read_csv(EXEC / "图像COX" / "00_performance.csv")
    cox_coef = pd.read_csv(EXEC / "图像COX" / "00_coefficients.csv")
    jm_perf = pd.read_csv(EXEC / "图像_JM" / "0827_JM_cindex.csv")
    jm_alpha = pd.read_csv(EXEC / "图像_JM" / "05_关联参数.csv")
    rsf = pd.read_csv(EXEC / "图像_RSF" / "00_performance.csv")
    rsflc = pd.read_csv(EXEC / "图像_RSFLC" / "00_performance.csv")
    return cox_perf, cox_coef, jm_perf, jm_alpha, rsf, rsflc


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    body = doc.element.body
    cox_perf, cox_coef, jm_perf, jm_alpha, rsf, rsflc = load_results()

    cox_test = cox_perf[cox_perf["dataset"] == "Test"].iloc[0]
    cox_train = cox_perf[cox_perf["dataset"] == "Train"].iloc[0]
    jm_test = jm_perf[jm_perf["dataset"] == "Test"].iloc[0]
    jm_train = jm_perf[jm_perf["dataset"] == "Train"].iloc[0]
    rsf_tr, rsf_te = rsf["training_set"].iloc[0], rsf["test_set"]
    rsf_map = dict(zip(rsf["metric"], rsf["test_set"]))
    rsf_tr_map = dict(zip(rsf["metric"], rsf["training_set"]))
    lc_map = dict(zip(rsflc["metric"], rsflc["test_set"]))
    lc_tr_map = dict(zip(rsflc["metric"], rsflc["training_set"]))

    # ---------- 删除 5.5.3 KM（标题+正文+图+题注）----------
    p_553 = find_body(body, lambda t: t.startswith("5.5.3"), last=True)
    p_56 = find_body(body, lambda t: t.startswith("5.6 真实数据分析结果讨论"), last=True)
    if p_553 is None or p_56 is None:
        raise RuntimeError("未定位 5.5.3 / 5.6")
    kids = body_children(body)
    i0, i1 = kids.index(p_553), kids.index(p_56)
    for el in kids[i0:i1]:
        body.remove(el)
    print(f"deleted 5.5.3 block n={i1 - i0}")

    # ---------- 全局替换（插入新表前把旧表5-5改为表5-7）----------
    mapping = {
        "5.5 校准、决策曲线与风险分层": "5.5 校准与决策曲线分析",
        "表5-5": "表5-7",
    }
    for p in doc.paragraphs:
        t = p.text
        nt = t
        for a, b in mapping.items():
            nt = nt.replace(a, b)
        if nt != t:
            set_text(p, nt)

    # ---------- 样式源 ----------
    style_h3 = style_body = style_cap = style_note = style_toc3 = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "5.3.3 最优超参数下的模型性能比较" and p.style.name == "Heading 3":
            style_h3 = p
        if t.startswith("采用第5.3.2节选定的最优超参数") and style_body is None:
            style_body = p
        if t.startswith("表5-4 独立验证集") and p.style.name == "表格标题":
            style_cap = p
        if t.startswith("注：RSF的Brier") or t.startswith("注：选参准则"):
            if "Brier" in t:
                style_note = p
        if t.startswith("5.3.3 最优超参数下的模型性能比较") and p.style.name.startswith("toc"):
            style_toc3 = p
    if not all([style_h3, style_body, style_cap, style_note, style_toc3]):
        raise RuntimeError(
            f"style missing h3={style_h3 is not None} body={style_body is not None} "
            f"cap={style_cap is not None} note={style_note is not None} toc={style_toc3 is not None}"
        )

    # TOC: 5.3.4 / 5.3.5
    toc_533 = None
    for p in doc.paragraphs[:90]:
        if p.text.strip().startswith("5.3.3 最优超参数下的模型性能比较") and p.style.name.startswith("toc"):
            toc_533 = p
            break
    el = toc_533._p
    el = clone_p_after(el, style_toc3, "5.3.5 JM关联参数与MCMC诊断")
    el = clone_p_after(toc_533._p, style_toc3, "5.3.4 COX比例风险模型的风险比")
    print("toc inserted 5.3.4/5.3.5")

    # ---------- 表5-4：四模型 ----------
    tbl54 = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("表5-4 独立验证集28天预测性能"):
            # next table in body after this para
            break
    # find by first-row header 模型 + 最优参数 or already changed
    for tbl in doc.tables:
        header_join = " | ".join(c.text.strip() for c in tbl.rows[0].cells)
        if "OOB C-index" in header_join or "测试集 C-index" in header_join:
            tbl54 = tbl
            break
    if tbl54 is None:
        raise RuntimeError("未找到表5-4")
    template_tbl_el = copy.deepcopy(tbl54._tbl)
    tbl_parent = tbl54._parent

    headers54 = [
        "模型",
        "模型设定",
        "训练集 C-index",
        "验证集 C-index",
        "28天AUC",
        "28天Brier",
        "IBS 0–28",
    ]
    rows54 = [
        [
            "COX",
            "8个基线 + 第1天20个截面；无需调参",
            f"{cox_train['cindex']:.4f}",
            f"{cox_test['cindex']:.4f}",
            [("缺失（待补）", True)],
            f"{cox_test['brier_28']:.4f}",
            f"{cox_test['ibs']:.4f}",
        ],
        [
            "JM",
            "纵向GCS/SOFA/CNS + 生存子模型",
            f"{jm_train['cindex']:.4f}",
            [("0.8472（待核）", True)],
            [("1.0000（不合理，待核）", True)],
            f"{float(jm_test['brier_28_ipcw']):.4f}",
            f"{float(jm_test['ibs_0_28_ipcw']):.4f}",
        ],
        [
            "RSF",
            "ntree=500, mtry=3, nodesize=10, nsplit=10",
            f"{float(rsf_tr_map['C-index']):.4f}",
            f"{float(rsf_map['C-index']):.4f}",
            f"{float(rsf_map['28-day AUC']):.4f}",
            f"{float(rsf_map['28-day Brier (IPCW)']):.4f}",
            f"{float(rsf_map['IBS 0-28 days (IPCW)']):.4f}",
        ],
        [
            "RSFLC",
            "ntree=200, mtry=3, nodesize=1, t0=5天",
            f"{float(lc_tr_map['C-index']):.4f}",
            f"{float(lc_map['C-index']):.4f}",
            f"{float(lc_map['28-day AUC']):.4f}",
            f"{float(lc_map['28-day Brier (IPCW)']):.4f}",
            f"{float(lc_map['IBS 0-28 days (IPCW)']):.4f}",
        ],
    ]
    rebuild_table(tbl54, headers54, rows54)
    print("rebuilt table 5-4")

    # ---------- 5.3.4 / 5.3.5 插在表5-4注之后、5.4之前 ----------
    p_note54 = find_body(body, lambda t: t.startswith("注：") and "Brier" in t, last=True)
    # 表5-4注可能已被后面逻辑改写前仍是旧文；用更稳的：紧邻5.4之前的注
    p_54head = find_body(body, lambda t: t.startswith("5.4 模型解释"), last=True)
    if p_54head is None:
        raise RuntimeError("未找到5.4")
    # 找 5.4 前一个有文本的段落作为注
    kids = body_children(body)
    i_54 = kids.index(p_54head)
    p_note54 = None
    for el in reversed(kids[:i_54]):
        if el.tag == qn("w:p") and p_text(el).startswith("注："):
            p_note54 = el
            break
    if p_note54 is None:
        raise RuntimeError("未找到表5-4表注")

    # COX 精简表
    coef = cox_coef.set_index("Variable")
    cox_rows_spec = [
        ("mechvent1", "机械通气", 2),
        ("electivesurgery1", "择期手术", 2),
        ("age", "年龄", 3),
        ("charlson_comorbidity_index", "Charlson合并症指数", 3),
        ("apsiii", "APSIII", 3),
        ("sapsii", "SAPSII", 3),
        ("oasis", "OASIS", 3),
        ("gcs", "GCS（第1天）", 3),
        ("glucose", "血糖（第1天）", 3),
        ("hemoglobin", "血红蛋白（第1天）", 3),
        ("ph", "pH（第1天）", 2),
    ]
    cox_table_rows = []
    for key, label, nd in cox_rows_spec:
        r = coef.loc[key]
        hr = fmt_hr(r["HR_95CI_Lower"], r["HR_95CI_Upper"], r["HR"], nd=nd)
        pv = fmt_p(float(r["P_value"]))
        if key == "ph":
            cox_table_rows.append(
                [
                    label,
                    [(hr + "（尺度异常，待核）", True)],
                    pv,
                ]
            )
        else:
            cox_table_rows.append([label, hr, pv])

    # JM 关联表
    alpha = jm_alpha.set_index("parameter")
    jm_labels = [
        ("value(gcs)", "GCS（当前值关联）"),
        ("value(sofa_24hours)", "SOFA总分（当前值关联）"),
        ("value(cns_24hours)", "SOFA神经系统（当前值关联）"),
    ]
    jm_table_rows = []
    for key, label in jm_labels:
        r = alpha.loc[key]
        hr = fmt_hr(r["hazard_ratio_low"], r["hazard_ratio_high"], r["hazard_ratio"], nd=2)
        rhat = float(r["Rhat"])
        mean = f"{float(r['posterior_mean']):.3f}"
        rhat_cell = [(f"{rhat:.3f}（>1.1，待核）", True)] if rhat > 1.1 else f"{rhat:.3f}"
        jm_table_rows.append([label, mean, hr, rhat_cell])

    # 先把 5.3.5 倒序插入，再插 5.3.4，使最终顺序为 4→5
    # 实际用正向：不断 addnext 在最新节点后
    el = p_note54

    el = clone_p_after(el, style_h3, "5.3.4 COX比例风险模型的风险比")
    el = clone_p_after(
        el,
        style_body,
        segments=[
            (
                "COX采用与RSF对齐的变量定义：8个基线协变量加上20条轨迹在第1天的截面，于全部训练集拟合比例风险模型后在独立验证集评价。"
                "验证集C-index为0.7740，与RSF接近；",
                False,
            ),
            ("28天AUC尚未由当前结果脚本输出，表5-4中以红色标注为“缺失（待补）”。", True),
            (
                "表5-5列出对临床解释较关键的协变量风险比。机械通气的死亡风险约为未通气者的2.92倍（95%CI 2.40–3.56）；"
                "择期手术与较低风险相关；年龄、Charlson合并症指数、APSIII与SAPSII的方向与临床认识一致；"
                "第1天GCS升高对应风险下降。OASIS在本拟合中未达到常规显著性水平。",
                False,
            ),
            (
                "血气pH的点估计HR约为36.97，可信区间跨越数个数量级，明显受到“每增加1个pH单位”这一不恰当尺度的影响，表中以红色标出，不宜按字面解释。",
                True,
            ),
            ("完整28个变量的系数表见建模输出，正文不逐一罗列。", False),
        ],
    )
    el = clone_p_after(el, style_cap, "表5-5 COX模型主要协变量风险比（训练集拟合）")
    el = insert_table_after(
        el,
        template_tbl_el,
        tbl_parent,
        ["变量", "HR（95%CI）", "P"],
        cox_table_rows,
    )
    el = clone_p_after(
        el,
        style_note,
        segments=[
            (
                "注：HR由回归系数指数变换得到。表中仅保留基线核心变量及部分显著实验室截面；",
                False,
            ),
            ("pH一行因回归尺度不合理以红色标注，待改用中心化或按0.01单位重新报告。", True),
            ("未列入的第1天截面系数见图像COX/00_coefficients.csv。", False),
        ],
    )

    el = clone_p_after(el, style_h3, "5.3.5 JM关联参数与MCMC诊断")
    el = clone_p_after(
        el,
        style_body,
        segments=[
            (
                "JM将GCS、SOFA总分及SOFA神经系统亚分作为纵向过程，与含8个临床基线和其余第1天截面的生存子模型通过当前值关联相连。"
                "验证集分析人数为1,699（事件311），少于COX的1,781（事件378）。",
                False,
            ),
            (
                "验证集C-index为0.8472，高于其余模型，但同一风险过程得到的28天AUC为1.0000，在临床队列中几乎不可能成立，故表5-4以红色标注；本节不以该C-index作为“JM全面优于森林模型”的证据。",
                True,
            ),
            (
                "表5-6给出三个关联参数：GCS潜在水平升高与死亡风险下降相关（HR 0.63），SOFA升高与风险上升相关（HR 1.23），神经系统亚分方向与GCS一致。生存子模型中机械通气、合并症指数与年龄的方向与COX相同。",
                False,
            ),
            (
                "MCMC为3条链、总迭代1500、预烧500。诊断摘要显示43个参数中13个Rhat>1.1，最小ESS约为10.7，ESS<100者10个，收敛不充分；GCS与神经系统关联参数的Rhat亦高于1.1（表5-6红色）。JM目前宜作为机制解释的参考，预测数字须在加长抽样并核对风险过程口径后再纳入终评排序。",
                True,
            ),
        ],
    )
    el = clone_p_after(el, style_cap, "表5-6 JM当前值关联参数（后验摘要）")
    el = insert_table_after(
        el,
        template_tbl_el,
        tbl_parent,
        ["纵向标记", "后验均数", "HR（95%CI）", "Rhat"],
        jm_table_rows,
    )
    el = clone_p_after(
        el,
        style_note,
        segments=[
            (
                "注：HR表示纵向指标潜在真实水平增加1个原始计量单位时死亡瞬时风险的相对变化。",
                False,
            ),
            ("Rhat>1.1的单元格以红色标注，提示收敛不足，数字仅供核验，不作为最终推断。", True),
            ("轨迹图、DIC/WAIC与全部后验参数见图像_JM输出，正文从略。", False),
        ],
    )
    print("inserted 5.3.4 and 5.3.5")

    # ---------- 重写关键段落 ----------
    replacements: list[tuple[str, list[tuple[str, bool]] | str]] = []

    replacements.append(
        (
            "本章利用MIMIC-IV数据库的脑卒中住院队列",
            (
                "本章利用MIMIC-IV数据库的脑卒中住院队列（含缺血性与出血性等亚型），对四种生存模型做真实的分析。"
                "分析流程分为五个部分，即（1）研究队列的详细描述和临床特征总结；（2）多源纵向数据预处理和融合；"
                "（3）在训练集内五折交叉验证中完成超参数选择，再以最优超参数于全训练集重拟合并在独立验证集评价预测性能；"
                "（4）在同一验证集上，对最优RSF与RSFLC模型进行置换变量重要性（VIMP）与SHAP解释；"
                "（5）基于28天死亡风险给出校准图与决策曲线分析。"
                "四种模型为COX、JM、RSF与RSFLC。"
            ),
        )
    )
    replacements.append(
        (
            "就具体模型设定而言：COX模型通常无需网格调参",
            (
                "就具体模型设定而言：COX模型通常无需网格调参，采用8个基线协变量与20条轨迹的第1天截面，可直接在训练集拟合后于验证集评价；"
                "JM模型将GCS、SOFA总分及SOFA神经系统亚分作为纵向过程，纵向子模型采用线性混合效应，生存子模型采用比例风险形式，并通过当前值关联连接纵向与生存部分；"
                "RSF在训练集内五折交叉验证中对树数、每次分裂候选变量数、终端节点最小样本等超参数进行网格搜索，再以最优组合在全训练集上重拟合；"
                "RSFLC为landmark动态随机森林，在时刻t0=5天纳入纵向轨迹与固定协变量，同样经训练集内五折交叉验证选定超参数后，在全训练集重拟合并于验证集评价。"
                "未死亡病例在第28天予以行政删失，以与短期预后终点对齐。"
            ),
        )
    )
    replacements.append(
        (
            "在完成验证集预测后，进一步对机器学习模型（RSF、RSFLC）开展置换变量重要性",
            (
                "在完成验证集预测后，进一步对机器学习模型（RSF、RSFLC）开展置换变量重要性（VIMP）与SHAP分析，"
                "以刻画在既定模型与评价协议下各协变量对28天死亡风险的相对贡献。"
                "该类解释性分析建立在上述最优模型及其验证集预测之上，与性能终评共用同一划分，而不另设并行评价协议。"
                "需要说明的是：表5-3给出RSF与RSFLC的超参数选择结果；"
                "表5-4给出四种模型在独立验证集上的预测性能终评，并作为第5.4–5.5节解释与临床效用分析的共同对照；"
                "表5-5与表5-6分别给出COX风险比与JM关联参数。"
                "不以训练集内交叉验证的选参折性能替代验证集终评数字。"
            ),
        )
    )
    replacements.append(
        (
            "采用第5.3.2节选定的最优超参数（COX与JM按既定模型设定、无需网格调参）",
            (
                "采用第5.3.2节选定的最优超参数（COX与JM按既定模型设定、无需网格调参），"
                "在全部训练集上重新拟合各模型，并在独立验证集上计算C-index、28天AUC、Brier Score与IBS等指标，"
                "作为本框架下的模型预测性能终评。该步骤不再重复使用训练集内五折交叉验证进行模型间主比较，"
                "以避免与超参数选择过程角色重叠，并与第5.3.1节“独立验证集评价”的设定保持一致。"
            ),
        )
    )

    for p in doc.paragraphs:
        t = p.text
        for prefix, new in replacements:
            if t.startswith(prefix):
                if isinstance(new, str):
                    set_text(p, new)
                else:
                    apply_segments(p, new)
                break

    # 表5-4导语（含红字）
    for p in doc.paragraphs:
        if p.text.startswith("表5-4汇总独立验证集上的主要预测性能"):
            apply_segments(
                p,
                [
                    (
                        "表5-4汇总四种模型在独立验证集上的主要预测性能，并同时给出训练集C-index，供第5.4–5.5节解释与临床效用分析对照。"
                        "各模型纳入的信息并不完全相同（COX与RSF以基线及第1天截面为主，RSFLC在t0=5天纳入纵向轨迹，JM以三条纵向过程连接生存子模型），",
                        False,
                    ),
                    (
                        "且分析人数不完全一致：COX训练/验证为4,390/1,781（事件901/378），JM为4,163/1,699（事件727/311）；RSF与RSFLC人数未在性能表中写出，待补。",
                        True,
                    ),
                    (
                        "比较以验证集C-index为主，Brier与IBS结合表注解读。",
                        False,
                    ),
                    (
                        "表中红色字体标出目前缺失或明显不合理、需要替换后再作排序的数字。",
                        True,
                    ),
                ],
            )

    # 表5-4题注
    for p in doc.paragraphs:
        if p.text.startswith("表5-4 独立验证集28天预测性能"):
            set_text(p, "表5-4 独立验证集28天预测性能（70%/30%划分；四种模型终评）")

    # 表5-4表注（含红字）——可能有两条以“注：”开头且含Brier的，取含IPCW或含RSF的那条且在5.3.3附近
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("注：RSF的Brier") or (
            t.startswith("注：") and "RSFLC为landmark" in t
        ):
            apply_segments(
                p,
                [
                    (
                        "注：Brier与IBS均为0–28天IPCW（RSF、RSFLC取自执行目录最新performance表；COX、JM取自对应结果目录）。"
                        "RSF训练集C-index为OOB估计，其余模型为表观C-index，不宜直接横比训练集一列。",
                        False,
                    ),
                    (
                        "COX的28天AUC缺失；JM的28天AUC=1.0000明显不合理，其验证集C-index与该AUC共用event-process风险，一并待核，故表5-4不以JM为性能最优结论。",
                        True,
                    ),
                    (
                        "RSF/RSFLC数字已按图像_RSF与图像_RSFLC的performance表更新，与第5.4节SHAP所用拟合是否为同一次运行待核。",
                        True,
                    ),
                ],
            )
            break

    # 5.4.2 数字
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("RSFLC为landmark时刻t0=5天的动态随机森林") and "测试集C-index为0.7736" in t:
            set_text(
                p,
                t.replace("测试集C-index为0.7736，28天AUC为0.8021。", "测试集C-index为0.7624，28天AUC为0.7781。"),
            )

    # 5.5 导语
    for p in doc.paragraphs:
        if p.text.startswith("区分度只说明模型能否把先发生事件的患者排在前面"):
            set_text(
                p,
                "区分度只说明模型能否把先发生事件的患者排在前面。临床使用还取决于预测概率是否校准、在不同决策阈值下是否带来净获益。"
                "以下分析均在独立验证集、28天结局上进行，与表5-4终评及第5.4节解释性分析共用同一划分；"
                "目前仅对已输出校准图与决策曲线的RSF、RSFLC报告，COX与JM尚无对应图形。",
            )

    # 5.6
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("真实数据中模型性能排序"):
            apply_segments(
                p,
                [
                    (
                        "若只看目前可用的验证集C-index点估计，RSF为0.7958，COX为0.7740，RSFLC为0.7624。",
                        False,
                    ),
                    (
                        "JM的验证集C-index为0.8472，但28天AUC为1.0000，明显不合理，故不能把“RSF＞RSFLC＞COX＞JM”或“JM最优”作为本章结论；在AUC口径修好之前，性能排序只在COX、RSF与RSFLC之间讨论。",
                        True,
                    ),
                    (
                        "在这三者中，RSF仍具有最高的验证集区分度，COX与之接近，RSFLC略低。真实数据中模型差距小于模拟研究，说明额外噪声会压缩方法之间的差异。",
                        False,
                    ),
                ],
            )
        if t.startswith("第二，JM模型在真实数据中的区分度仍弱于机器学习模型"):
            apply_segments(
                p,
                [
                    (
                        "第二，JM在本队列中的机制解释仍有价值：关联参数显示GCS潜在水平升高对应死亡风险下降、SOFA升高对应风险上升，与临床方向一致；生存子模型中机械通气等系数也与COX同向。",
                        False,
                    ),
                    (
                        "但其预测数字目前不能进入终评排序：除AUC=1.00外，MCMC诊断中13个参数Rhat>1.1、最小ESS约10.7，收敛不充分。旧稿曾报告JM的交叉验证AUC=0.6563，与本次独立验证集结果不是同一评价协议，亦不得混用。",
                        True,
                    ),
                    (
                        "即便将来预测口径修正，仍需警惕线性纵向轨迹假设在GCS、SOFA等指标上被违反，以及纵向子模型与生存子模型设定错误会同时损害联合模型。",
                        False,
                    ),
                ],
            )
        if t.startswith("第五，测试集28天校准图显示"):
            set_text(
                p,
                "第五，测试集28天校准图显示RSF与RSFLC在低风险端拟合较好、在高风险端存在高估，故预测概率更宜用于排序，而非直接报告绝对死亡概率。"
                "尽管如此，决策曲线在约0.05–0.50的阈值范围内优于默认策略，说明两模型仍具有按风险采取差异化处理的临床参考价值。"
                "COX与JM尚未输出校准图与决策曲线，本节不对其绝对概率作同等解读。",
            )

    # 5.7
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("本章以MIMIC-IV数据库中6299次"):
            set_text(
                p,
                "本章以MIMIC-IV数据库中6299次脑卒中住院记录（对应6092名患者；含缺血性与出血性等亚型）的真实临床数据为基础，"
                "采用“训练集内五折交叉验证调参—全训练集重拟合—独立验证集评价”的框架，"
                "结合验证集上的VIMP、SHAP、校准和决策曲线，"
                "对四种生存分析模型（COX、JM、RSF、RSFLC）的预测性能与可解释性做评价。",
            )
        if t.startswith("在独立验证集终评下，机器学习模型整体具有更好的区分度表现"):
            apply_segments(
                p,
                [
                    (
                        "在独立验证集终评下，RSF的验证集C-index为0.7958、28天AUC为0.8222；"
                        "RSFLC的验证集C-index为0.7624、AUC为0.7781；COX的验证集C-index为0.7740。",
                        False,
                    ),
                    (
                        "COX的28天AUC缺失；JM的验证集C-index为0.8472、AUC为1.0000，后者明显不合理，二者均以红色标于表5-4，不作为本章排序依据。",
                        True,
                    ),
                    (
                        "该结果与第5.3.2节选定的最优超参数相对应。COX风险比与JM关联参数见第5.3.4–5.3.5节，变量贡献与临床效用见第5.4节与第5.5节。",
                        False,
                    ),
                ],
            )
        if t.startswith("在独立验证集与28天结局下，RSF的验证集C-index为0.7885"):
            apply_segments(
                p,
                [
                    (
                        "在独立验证集与28天结局下，目前可用于排序的模型为RSF（C-index 0.7958，AUC 0.8222）、"
                        "COX（C-index 0.7740）与RSFLC（C-index 0.7624，AUC 0.7781）。"
                        "VIMP与SHAP表明，机械通气和SAPSII是最稳定的风险贡献因素，其次为APSIII、OASIS、年龄和CCI；"
                        "RSFLC中GCS轨迹的置换重要性为0。校准图提示高危端存在高估，但决策曲线在较宽阈值范围内优于默认策略。"
                        "表5-3用于超参数选择过程说明，表5-4用于外样本预测终评；二者分工不同，引用时按表注区分。",
                        False,
                    ),
                    (
                        "JM与COX中以红色标注的缺失或异常指标须在补算、核对风险过程并改善MCMC收敛后再写入最终排序。",
                        True,
                    ),
                    (
                        "结合前面章节的模拟结果，本文为临床生存预测任务中的模型选择提供了可对照的证据。",
                        False,
                    ),
                ],
            )

    doc.save(str(OUT))
    print("saved", OUT)

    # ---------- 检查 ----------
    doc2 = Document(str(OUT))
    print("==== structure ====")
    seen = []
    for i, p in enumerate(doc2.paragraphs):
        t = p.text.strip()
        if t.startswith(("5.3.", "5.4", "5.5", "5.6", "5.7", "表5-4", "表5-5", "表5-6", "表5-7", "图5-")):
            if i < 90 or i > 200:
                print(f"{i} [{p.style.name}] {t[:90]}")
                seen.append(t[:40])
    texts = "\n".join(p.text for p in doc2.paragraphs)
    print("has 5.5.3", "5.5.3" in texts)
    print("has 图5-8", "图5-8" in texts)
    print("has Kaplan–Meier风险分层 heading", "5.5.3 Kaplan" in texts)
    red_n = 0
    for p in doc2.paragraphs:
        for r in p.runs:
            rPr = r._r.rPr
            if rPr is None:
                continue
            c = rPr.find(qn("w:color"))
            if c is not None and (c.get(qn("w:val")) or "").upper() == RED_VAL:
                red_n += 1
    for tbl in doc2.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        rPr = r._r.rPr
                        if rPr is None:
                            continue
                        c = rPr.find(qn("w:color"))
                        if c is not None and (c.get(qn("w:val")) or "").upper() == RED_VAL:
                            red_n += 1
    print("red runs", red_n)
    # table 5-4 preview
    for tbl in doc2.tables:
        h0 = [c.text.strip()[:12] for c in tbl.rows[0].cells]
        if h0 and h0[0] == "模型" and "28天AUC" in "".join(c.text for c in tbl.rows[0].cells):
            print("T54", h0)
            for row in tbl.rows[1:]:
                print(" ", [c.text.strip()[:28] for c in row.cells])
            break


if __name__ == "__main__":
    main()
