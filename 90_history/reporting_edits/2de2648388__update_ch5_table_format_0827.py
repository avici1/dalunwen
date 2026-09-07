# -*- coding: utf-8 -*-
"""统一表5-1~5-7格式（以5-6/5-7为准），并向表5-1追加性别/BMI/分型。"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

ROOT = Path(r"F:\文章_大论文")
DOCX_IN = ROOT / r"0722\大论文版本\0826_1_基于机器学习的生存分析方法探讨_第五章RSF_RSFLC图更新.docx"
DOCX_OUT = ROOT / r"0722\大论文版本\0827_1_基于机器学习的生存分析方法探讨_第五章表格式与表5-1更新.docx"
CSV_NEW = ROOT / r"0722\实例研究代码\stroke_baseline_knn_0824_fold.csv"

# 5-6/5-7 参考样式
HEADER_FILL = "D9E2F3"
FONT_ASCII = "Times New Roman"
FONT_EA = "宋体"
FONT_SIZE_PT = 10.5


def set_run_font(run, bold: bool | None = None):
    run.font.name = FONT_ASCII
    run.font.size = Pt(FONT_SIZE_PT)
    if bold is not None:
        run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), FONT_ASCII)
    rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    rFonts.set(qn("w:eastAsia"), FONT_EA)
    rFonts.set(qn("w:cs"), FONT_ASCII)


def set_cell_text(cell, text: str, *, bold: bool = False, align: str = "center"):
    # clear existing paragraphs content
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        set_run_font(p0.runs[0], bold=bold)
        for r in p0.runs[1:]:
            r.text = ""
    else:
        run = p0.add_run(text)
        set_run_font(run, bold=bold)

    # paragraph alignment
    pPr = p0._p.get_or_add_pPr()
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), align)

    # vertical center
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), "center")


def set_cell_shading(cell, fill: str | None):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tcPr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_tbl_border_el(borders, tag: str, val="single", color="000000", sz="4"):
    el = borders.find(qn(f"w:{tag}"))
    if el is None:
        el = OxmlElement(f"w:{tag}")
        borders.append(el)
    el.set(qn("w:val"), val)
    el.set(qn("w:color"), color)
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "0")


def apply_ref_table_format(table):
    """把表格格式改成与表5-6/5-7一致。"""
    try:
        table.style = "Normal Table"
    except Exception:
        pass

    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    # borders black full grid (same as 5-6/5-7 explicit borders)
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        set_tbl_border_el(borders, edge)

    # width 100% (5000 pct), center
    def upsert(tag, attrs):
        el = tblPr.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tblPr.append(el)
        for k, v in attrs.items():
            el.set(qn(f"w:{k}"), v)
        return el

    upsert("tblW", {"w": "5000", "type": "pct"})
    upsert("jc", {"val": "center"})
    upsert("tblLayout", {"type": "autofit"})

    # cell margins like ref
    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        tblPr.append(mar)
    for side, w in [("top", "0"), ("left", "108"), ("bottom", "0"), ("right", "108")]:
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), w)
        el.set(qn("w:type"), "dxa")

    n_cols = len(table.columns)
    for ri, row in enumerate(table.rows):
        is_header = ri == 0
        for ci in range(n_cols):
            cell = row.cells[ci]
            text = cell.text.strip()
            # determine alignment: header all center; body col0 left else center
            if is_header:
                align = "center"
                bold = True
                set_cell_shading(cell, HEADER_FILL)
            else:
                align = "left" if ci == 0 else "center"
                # section-like rows (only first col filled / empty others) keep bold left
                others = [row.cells[j].text.strip() for j in range(1, n_cols)]
                is_section = text != "" and all(o in ("", "—", "-") for o in others) and ci == 0
                # also common section titles
                section_names = {
                    "队列规模",
                    "人口学与入院特征",
                    "入院严重度评分",
                    "结局（供参考）",
                    "卒中分型",
                    "纵向随访概览（每住院观测天数）",
                    "患者-日记录总数",
                }
                if text in section_names:
                    is_section = True
                bold = bool(is_section and ci == 0)
                set_cell_shading(cell, None)
            set_cell_text(cell, text, bold=bold, align=align)


def mean_sd(x: pd.Series, digits: int = 2) -> str:
    v = pd.to_numeric(x, errors="coerce").dropna()
    if len(v) == 0:
        return "—"
    return f"{v.mean():.{digits}f} ± {v.std(ddof=1):.{digits}f}"


def n_pct(mask_or_series, n_total: int | None = None) -> str:
    if isinstance(mask_or_series, pd.Series) and mask_or_series.dtype != bool:
        v = pd.to_numeric(mask_or_series, errors="coerce").dropna()
        k = int((v == 1).sum())
        n = len(v) if n_total is None else n_total
    else:
        m = mask_or_series.fillna(False) if hasattr(mask_or_series, "fillna") else mask_or_series
        k = int(np.sum(m))
        n = int(len(m)) if n_total is None else n_total
    if n == 0:
        return "—"
    return f"{k:,} ({100.0 * k / n:.1f})"


def find_ch5_tables(doc: Document):
    """按内容识别表5-1~5-7。"""
    mapping = {}
    for ti, tbl in enumerate(doc.tables):
        rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
        if not rows:
            continue
        flat = " | ".join("|".join(r) for r in rows)
        head = "|".join(rows[0])
        if (
            "年龄" in flat
            and "Charlson" in flat
            and "样本量" in head
            and "均值±标准差" in head
            and "有观测住院数" not in head
        ):
            mapping["5-1"] = ti
        elif "第1天（均值±SD）" in head or ("有观测住院数" in head and "最差值" in head):
            mapping["5-2"] = ti
        elif (
            "模型" in head
            and "AUC" in head
            and "C-index" in head
            and "Brier" in head
            and ("±" in flat)
            and ("0.6761" in flat or "0.7977" in flat)
        ):
            mapping["5-3"] = ti
        elif ("指标" in head) and ("χ" in flat or "自由度" in flat):
            mapping["5-4"] = ti
        elif "比较组" in head and ("0.062" in flat or "Wilcoxon" in flat or "COX vs RSF" in flat):
            mapping["5-5"] = ti
        elif "最优参数" in head and "28天AUC" in head:
            mapping["5-6"] = ti
        elif "RSF VIMP" in head and "RSFLC VIMP" in head:
            mapping["5-7"] = ti
    return mapping


def insert_rows_after(table, after_idx: int, new_rows: list[list[str]]):
    """在 after_idx 行之后插入若干行（复制该行结构）。"""
    tbl = table._tbl
    ref_tr = table.rows[after_idx]._tr
    for vals in new_rows:
        new_tr = copy.deepcopy(ref_tr)
        # clear text
        for t in new_tr.iter(qn("w:t")):
            t.text = ""
        ref_tr.addnext(new_tr)
        ref_tr = new_tr
        # fill via python-docx row wrapper: find the just-added row index
    # After all inserts, fill texts by scanning
    # Easier: re-find rows after insert position
    # Actually addnext inserts immediately after, so order of new_rows is preserved
    # Fill from after_idx+1
    for i, vals in enumerate(new_rows):
        row = table.rows[after_idx + 1 + i]
        for j, val in enumerate(vals):
            if j < len(row.cells):
                # temporary plain text; format applied later
                for p in row.cells[j].paragraphs:
                    for r in p.runs:
                        r.text = ""
                p0 = row.cells[j].paragraphs[0]
                if p0.runs:
                    p0.runs[0].text = val
                else:
                    p0.add_run(val)


def add_baseline_extra_vars(table):
    """向表5-1追加性别、BMI、分型。"""
    df = pd.read_csv(CSV_NEW)
    n = len(df)
    # find anchor row index: 年龄
    age_idx = None
    for i, row in enumerate(table.rows):
        if row.cells[0].text.strip().startswith("年龄"):
            age_idx = i
            break
    if age_idx is None:
        raise RuntimeError("表5-1未找到年龄行")

    # gender
    male_n = int((df["gender"].astype(str).str.upper() == "M").sum())
    male_pct = f"{male_n:,} ({100.0 * male_n / n:.1f})"
    # BMI
    bmi = pd.to_numeric(df["bmi"], errors="coerce")
    bmi_n = int(bmi.notna().sum())
    bmi_str = mean_sd(bmi)
    # stroke_type
    st = df["stroke_type"].fillna("缺失").astype(str)
    type_order = ["I63_脑梗死", "I61_脑出血", "I60_蛛网膜下腔出血", "混合型"]
    type_rows = []
    for lab in type_order:
        k = int((st == lab).sum())
        # nicer Chinese label
        nice = {
            "I63_脑梗死": "缺血性卒中（I63脑梗死），例(%)",
            "I61_脑出血": "脑出血（I61），例(%)",
            "I60_蛛网膜下腔出血": "蛛网膜下腔出血（I60），例(%)",
            "混合型": "混合型，例(%)",
        }[lab]
        type_rows.append([nice, f"{n:,}", f"{k:,} ({100.0 * k / n:.1f})"])

    extra = [
        ["男性，例(%)", f"{n:,}", male_pct],
        ["体重指数（BMI）", f"{bmi_n:,}", bmi_str],
        ["卒中分型", "", ""],
        *type_rows,
    ]
    insert_rows_after(table, age_idx, extra)
    print("表5-1 已追加: 男性 / BMI / 卒中分型")
    print("  男性", male_pct, "BMI", bmi_str, "n=", bmi_n)
    for r in type_rows:
        print(" ", r[0], r[2])


def main():
    shutil.copy2(DOCX_IN, DOCX_OUT)
    doc = Document(str(DOCX_OUT))
    mp = find_ch5_tables(doc)
    print("识别到的表:", mp)
    need = ["5-1", "5-2", "5-3", "5-4", "5-5", "5-6", "5-7"]
    missing = [k for k in need if k not in mp]
    if missing:
        raise RuntimeError(f"未完整识别表: missing={missing}, found={mp}")

    # 1) 先给表5-1加内容
    t51 = doc.tables[mp["5-1"]]
    add_baseline_extra_vars(t51)

    # 2) 对表5-1~5-7统一格式（含刚插入的行）
    # 重新识别以防行变化不影响 index
    mp2 = find_ch5_tables(doc)
    for key in need:
        ti = mp2[key]
        print(f"格式化 表{key} -> tables[{ti}]")
        apply_ref_table_format(doc.tables[ti])

    doc.save(str(DOCX_OUT))
    print("saved:", DOCX_OUT)

    # verify
    doc2 = Document(str(DOCX_OUT))
    mp3 = find_ch5_tables(doc2)
    t = doc2.tables[mp3["5-1"]]
    print("==== 表5-1 验证 ====")
    for r in t.rows:
        print(" | ".join(c.text.strip() for c in r.cells))
    # format spot check
    for key in ["5-1", "5-3", "5-6"]:
        tbl = doc2.tables[mp3[key]]
        cell = tbl.rows[0].cells[0]
        shd = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
        fill = shd.get(qn("w:fill")) if shd is not None else None
        run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else None
        print(
            f"表{key}: style={tbl.style.name}, header_fill={fill}, "
            f"font={run.font.name if run else None}, size={run.font.size.pt if run and run.font.size else None}"
        )


if __name__ == "__main__":
    main()
