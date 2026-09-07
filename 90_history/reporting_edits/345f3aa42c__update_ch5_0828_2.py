# -*- coding: utf-8 -*-
"""用 0828 晚间 JM 重跑结果更新 0828_1 第五章数字，另存 0828_2。"""
from __future__ import annotations

import importlib.util
import shutil
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(r"F:\文章_大论文")
SRC = ROOT / r"0722\大论文版本\0828_1_基于机器学习的生存分析方法探讨_第五章四模型终评与KM删除.docx"
OUT = ROOT / r"0722\大论文版本\0828_2_基于机器学习的生存分析方法探讨_第五章四模型终评数据更新.docx"
EXEC = ROOT / r"0722\实例研究代码\执行"
HELPER = ROOT / r"0722\_revise_ch5_0828_1.py"
RED_VAL = "FF0000"

spec = importlib.util.spec_from_file_location("ch5rev", HELPER)
mod = importlib.util.module_from_spec(spec)
sys.modules["ch5rev"] = mod
spec.loader.exec_module(mod)

apply_segments = mod.apply_segments
rebuild_table = mod.rebuild_table
fmt_hr = mod.fmt_hr


def is_red_run(run) -> bool:
    rPr = run._r.rPr
    if rPr is None:
        return False
    c = rPr.find(qn("w:color"))
    return c is not None and (c.get(qn("w:val")) or "").upper() == RED_VAL


def load_all():
    cox_perf = pd.read_csv(EXEC / "图像COX" / "00_performance.csv")
    jm_perf = pd.read_csv(EXEC / "图像_JM" / "0827_JM_cindex.csv")
    jm_alpha = pd.read_csv(EXEC / "图像_JM" / "05_关联参数.csv")
    jm_surv = pd.read_csv(EXEC / "图像_JM" / "04_生存子模型.csv")
    rsf = pd.read_csv(EXEC / "图像_RSF" / "00_performance.csv")
    rsflc = pd.read_csv(EXEC / "图像_RSFLC" / "00_performance.csv")
    post = pd.read_csv(EXEC / "图像_JM" / "03to06_全部后验参数.csv")
    return cox_perf, jm_perf, jm_alpha, jm_surv, rsf, rsflc, post


def mcmc_from_posterior(post: pd.DataFrame) -> dict:
    mask = ~post["parameter"].astype(str).str.startswith("frailty")
    sub = post[mask]
    rhat = pd.to_numeric(sub["Rhat"], errors="coerce")
    ess = pd.to_numeric(sub["ESS"], errors="coerce")
    return {
        "n_rhat": int(rhat.notna().sum()),
        "n_rhat_gt_11": int((rhat > 1.1).sum()),
        "ess_min": float(ess.min()),
        "n_ess_lt_100": int((ess < 100).sum()),
    }


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    cox_perf, jm_perf, jm_alpha, jm_surv, rsf, rsflc, post = load_all()
    mcmc = mcmc_from_posterior(post)

    cox_test = cox_perf[cox_perf["dataset"] == "Test"].iloc[0]
    cox_train = cox_perf[cox_perf["dataset"] == "Train"].iloc[0]
    jm_test = jm_perf[jm_perf["dataset"] == "Test"].iloc[0]
    jm_train = jm_perf[jm_perf["dataset"] == "Train"].iloc[0]
    rsf_map = dict(zip(rsf["metric"], rsf["test_set"]))
    rsf_tr_map = dict(zip(rsf["metric"], rsf["training_set"]))
    lc_map = dict(zip(rsflc["metric"], rsflc["test_set"]))
    lc_tr_map = dict(zip(rsflc["metric"], rsflc["training_set"]))

    jm_c_tr = float(jm_train["cindex"])
    jm_c_te = float(jm_test["cindex"])
    jm_auc_te = float(jm_test["auc_28"])
    jm_brier_te = float(jm_test["brier_28_ipcw"])
    jm_ibs_te = float(jm_test["ibs_0_28_ipcw"])
    cox_c_te = float(cox_test["cindex"])
    rsf_c = float(rsf_map["C-index"])
    rsf_auc = float(rsf_map["28-day AUC"])
    lc_c = float(lc_map["C-index"])
    lc_auc = float(lc_map["28-day AUC"])

    mech = jm_surv.set_index("parameter").loc["mechvent1"]
    mech_hr = f"{float(mech['hazard_ratio']):.2f}"

    # RSF/RSFLC 人数：性能表无 n，按与结果脚本相同的 group-fold 文件计数
    # RSF：基线 left_join 第1天截面，不去完整病例 → 4480/1819（966/408）
    # RSFLC：与有纵向记录者内连接 → 4390/1781（901/378）
    rsf_n = "4,480/1,819（事件966/408）"
    lc_n = "4,390/1,781（事件901/378）"
    cox_n = f"{int(cox_train['n']):,}/{int(cox_test['n']):,}（事件{int(cox_train['events'])}/{int(cox_test['events'])}）"
    jm_n = f"{int(jm_train['n']):,}/{int(jm_test['n']):,}（事件{int(jm_train['events'])}/{int(jm_test['events'])}）"

    # ---------- 表 5-4 ----------
    tbl54 = None
    for tbl in doc.tables:
        header_join = " | ".join(c.text.strip() for c in tbl.rows[0].cells)
        if "28天AUC" in header_join and header_join.startswith("模型"):
            tbl54 = tbl
            break
    if tbl54 is None:
        raise RuntimeError("未找到表5-4")

    headers54 = [
        "模型",
        "模型设定",
        "训练集 C-index",
        "验证集 C-index",
        "28天AUC",
        "28天Brier",
        "IBS 0–28",
    ]
    rows54 = [
        [
            "COX",
            "8个基线 + 第1天20个截面；无需调参",
            f"{float(cox_train['cindex']):.4f}",
            f"{cox_c_te:.4f}",
            [("缺失（待补）", True)],
            f"{float(cox_test['brier_28']):.4f}",
            f"{float(cox_test['ibs']):.4f}",
        ],
        [
            "JM",
            "纵向GCS/SOFA/CNS + 生存子模型",
            f"{jm_c_tr:.4f}",
            [(f"{jm_c_te:.4f}（待核）", True)],
            [(f"{jm_auc_te:.4f}（不合理，待核）", True)],
            f"{jm_brier_te:.4f}",
            f"{jm_ibs_te:.4f}",
        ],
        [
            "RSF",
            "ntree=500, mtry=3, nodesize=10, nsplit=10",
            f"{float(rsf_tr_map['C-index']):.4f}",
            f"{rsf_c:.4f}",
            f"{rsf_auc:.4f}",
            f"{float(rsf_map['28-day Brier (IPCW)']):.4f}",
            f"{float(rsf_map['IBS 0-28 days (IPCW)']):.4f}",
        ],
        [
            "RSFLC",
            "ntree=200, mtry=3, nodesize=1, t0=5天",
            f"{float(lc_tr_map['C-index']):.4f}",
            f"{lc_c:.4f}",
            f"{lc_auc:.4f}",
            f"{float(lc_map['28-day Brier (IPCW)']):.4f}",
            f"{float(lc_map['IBS 0-28 days (IPCW)']):.4f}",
        ],
    ]
    rebuild_table(tbl54, headers54, rows54)
    print("rebuilt table 5-4")

    # ---------- 表 5-6 ----------
    tbl56 = None
    for tbl in doc.tables:
        h0 = [c.text.strip() for c in tbl.rows[0].cells]
        if h0[:4] == ["纵向标记", "后验均数", "HR（95%CI）", "Rhat"]:
            tbl56 = tbl
            break
    if tbl56 is None:
        raise RuntimeError("未找到表5-6")

    alpha = jm_alpha.set_index("parameter")
    jm_labels = [
        ("value(gcs)", "GCS（当前值关联）"),
        ("value(sofa_24hours)", "SOFA总分（当前值关联）"),
        ("value(cns_24hours)", "SOFA神经系统（当前值关联）"),
    ]
    jm_table_rows = []
    alpha_hr_txt = {}
    for key, label in jm_labels:
        r = alpha.loc[key]
        hr = fmt_hr(r["hazard_ratio_low"], r["hazard_ratio_high"], r["hazard_ratio"], nd=2)
        rhat = float(r["Rhat"])
        mean = f"{float(r['posterior_mean']):.3f}"
        alpha_hr_txt[key] = hr
        rhat_cell = [(f"{rhat:.3f}（>1.1，待核）", True)] if rhat > 1.1 else f"{rhat:.3f}"
        jm_table_rows.append([label, mean, hr, rhat_cell])
    rebuild_table(
        tbl56,
        ["纵向标记", "后验均数", "HR（95%CI）", "Rhat"],
        jm_table_rows,
    )
    print("rebuilt table 5-6")
    print("alpha HR", alpha_hr_txt)
    print("mcmc", mcmc)

    gcs_hr = alpha_hr_txt["value(gcs)"]
    sofa_hr = alpha_hr_txt["value(sofa_24hours)"]
    cns_hr = alpha_hr_txt["value(cns_24hours)"]
    gcs_rhat = float(alpha.loc["value(gcs)", "Rhat"])
    sofa_rhat = float(alpha.loc["value(sofa_24hours)", "Rhat"])
    cns_rhat = float(alpha.loc["value(cns_24hours)", "Rhat"])

    # ---------- 段落 ----------
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("表5-4汇总四种模型在独立验证集上的主要预测性能"):
            apply_segments(
                p,
                [
                    (
                        "表5-4汇总四种模型在独立验证集上的主要预测性能，并同时给出训练集C-index，供第5.4–5.5节解释与临床效用分析对照。"
                        "各模型纳入的信息并不完全相同（COX与RSF以基线及第1天截面为主，RSFLC在t0=5天纳入纵向轨迹，JM以三条纵向过程连接生存子模型），"
                        f"分析人数亦不完全一致：COX训练/验证为{cox_n}；RSF为{rsf_n}，其中含缺失协变量并由随机生存森林插补；"
                        f"RSFLC为{lc_n}；JM预测指标在landmark=5仍存活者中计算，为{jm_n}。"
                        "比较以验证集C-index为主，Brier与IBS结合表注解读。",
                        False,
                    ),
                    (
                        "表中红色字体标出目前缺失或明显不合理、需要替换后再作排序的数字。",
                        True,
                    ),
                ],
            )
        elif t.startswith("注：Brier与IBS均为") or (
            t.startswith("注：") and "COX的28天AUC" in t
        ):
            apply_segments(
                p,
                [
                    (
                        "注：Brier与IBS均为IPCW（RSF、RSFLC取自执行目录最新performance表；COX、JM取自对应结果目录）。"
                        "RSF训练集C-index为OOB估计，其余模型为表观C-index，不宜直接横比训练集一列。"
                        "JM的C-index、AUC、Brier与IBS在landmark=5仍存活且有t≤5纵向记录的患者中计算，与COX/RSF的入组时点预测不是同一风险队列。",
                        False,
                    ),
                    (
                        f"COX的28天AUC缺失；JM的28天AUC={jm_auc_te:.4f}明显不合理，其验证集C-index与该AUC共用event-process风险，一并待核，故表5-4不以JM为性能最优结论。",
                        True,
                    ),
                    (
                        "RSF/RSFLC数字已按图像_RSF与图像_RSFLC的performance表更新，与第5.4节SHAP所用拟合是否为同一次运行待核。",
                        True,
                    ),
                ],
            )
        elif t.startswith("JM将GCS、SOFA总分及SOFA神经系统亚分作为纵向过程"):
            apply_segments(
                p,
                [
                    (
                        "JM将GCS、SOFA总分及SOFA神经系统亚分作为纵向过程，与含8个临床基线和其余第1天截面的生存子模型通过当前值关联相连。"
                        f"验证集预测指标在landmark=5仍存活者中计算，人数为{int(jm_test['n']):,}（事件{int(jm_test['events'])}），"
                        f"少于COX的{int(cox_test['n']):,}（事件{int(cox_test['events'])}）。",
                        False,
                    ),
                    (
                        f"验证集C-index为{jm_c_te:.4f}，高于其余模型，但同一风险过程得到的28天AUC为{jm_auc_te:.4f}，"
                        "在临床队列中几乎不可能成立，故表5-4以红色标注；本节不以该C-index作为“JM全面优于森林模型”的证据。",
                        True,
                    ),
                    (
                        f"表5-6给出三个关联参数：GCS当前值关联HR为{gcs_hr}，SOFA总分为{sofa_hr}，神经系统亚分为{cns_hr}。"
                        f"生存子模型中机械通气（HR约{mech_hr}）、合并症指数与年龄的方向与COX相同。",
                        False,
                    ),
                    (
                        f"GCS与神经系统关联的95%可信区间均包含1，不能写成“GCS升高对应死亡风险下降”。"
                        f"三个关联参数的Rhat分别为{gcs_rhat:.3f}、{sofa_rhat:.3f}与{cns_rhat:.3f}，均超过1.1（表5-6红色）。"
                        "MCMC为3条链、总迭代1500、预烧500。"
                        f"排除个体frailty后，{mcmc['n_rhat']}个有Rhat的参数中{mcmc['n_rhat_gt_11']}个Rhat>1.1，"
                        f"最小ESS约为{mcmc['ess_min']:.1f}，ESS<100者{mcmc['n_ess_lt_100']}个，收敛不充分。"
                        "JM目前宜作为机制方向的参考，预测数字与关联点估计均须在加长抽样并核对风险过程口径后再纳入终评排序。",
                        True,
                    ),
                ],
            )
        elif t.startswith("若只看目前可用的验证集C-index点估计"):
            apply_segments(
                p,
                [
                    (
                        f"若只看目前可用的验证集C-index点估计，RSF为{rsf_c:.4f}，COX为{cox_c_te:.4f}，RSFLC为{lc_c:.4f}。",
                        False,
                    ),
                    (
                        f"JM的验证集C-index为{jm_c_te:.4f}，但28天AUC为{jm_auc_te:.4f}，明显不合理，"
                        "故不能把“RSF＞RSFLC＞COX＞JM”或“JM最优”作为本章结论；"
                        "在AUC口径修好之前，性能排序只在COX、RSF与RSFLC之间讨论。",
                        True,
                    ),
                    (
                        "在这三者中，RSF仍具有最高的验证集区分度，COX与之接近，RSFLC略低。"
                        "真实数据中模型差距小于模拟研究，说明额外噪声会压缩方法之间的差异。",
                        False,
                    ),
                ],
            )
        elif t.startswith("第二，JM在本队列中的机制解释仍有价值"):
            apply_segments(
                p,
                [
                    (
                        f"第二，JM生存子模型中机械通气（HR约{mech_hr}）、合并症指数与年龄的方向与COX相同，可作为机制对照。"
                        f"关联参数方面，本轮后验中SOFA总分HR为{sofa_hr}，方向仍为风险上升；"
                        f"GCS为{gcs_hr}，神经系统亚分为{cns_hr}，二者区间包含1。",
                        False,
                    ),
                    (
                        f"因此不能再把“GCS潜在水平升高对应死亡风险下降”写成已证实结论。"
                        f"三个关联参数Rhat均>1.1；排除frailty后{mcmc['n_rhat_gt_11']}个参数Rhat>1.1、最小ESS约{mcmc['ess_min']:.1f}，收敛不充分。"
                        f"预测数字目前不能进入终评排序：除AUC={jm_auc_te:.4f}外，评价队列为landmark=5幸存者，与其余模型入组时点不同。"
                        "旧稿曾报告JM的交叉验证AUC=0.6563，与本次独立验证集结果不是同一评价协议，亦不得混用。",
                        True,
                    ),
                    (
                        "即便将来预测口径修正，仍需警惕线性纵向轨迹假设在GCS、SOFA等指标上被违反，"
                        "以及纵向子模型与生存子模型设定错误会同时损害联合模型。",
                        False,
                    ),
                ],
            )
        elif t.startswith("在独立验证集终评下，RSF的验证集C-index为"):
            apply_segments(
                p,
                [
                    (
                        f"在独立验证集终评下，RSF的验证集C-index为{rsf_c:.4f}、28天AUC为{rsf_auc:.4f}；"
                        f"RSFLC的验证集C-index为{lc_c:.4f}、AUC为{lc_auc:.4f}；"
                        f"COX的验证集C-index为{cox_c_te:.4f}。",
                        False,
                    ),
                    (
                        f"COX的28天AUC缺失；JM的验证集C-index为{jm_c_te:.4f}、AUC为{jm_auc_te:.4f}，"
                        "后者明显不合理，二者均以红色标于表5-4，不作为本章排序依据。",
                        True,
                    ),
                    (
                        "该结果与第5.3.2节选定的最优超参数相对应。"
                        "COX风险比与JM关联参数见第5.3.4–5.3.5节，变量贡献与临床效用见第5.4节与第5.5节。",
                        False,
                    ),
                ],
            )
        elif t.startswith("在独立验证集与28天结局下，目前可用于排序的模型为"):
            apply_segments(
                p,
                [
                    (
                        f"在独立验证集与28天结局下，目前可用于排序的模型为RSF（C-index {rsf_c:.4f}，AUC {rsf_auc:.4f}）、"
                        f"COX（C-index {cox_c_te:.4f}）与RSFLC（C-index {lc_c:.4f}，AUC {lc_auc:.4f}）。"
                        "VIMP与SHAP表明，机械通气和SAPSII是最稳定的风险贡献因素，其次为APSIII、OASIS、年龄和CCI；"
                        "RSFLC中GCS轨迹的置换重要性为0。校准图提示高危端存在高估，但决策曲线在较宽阈值范围内优于默认策略。"
                        "表5-3用于超参数选择过程说明，表5-4用于外样本预测终评；二者分工不同，引用时按表注区分。",
                        False,
                    ),
                    (
                        "JM与COX中以红色标注的缺失或异常指标须在补算、核对风险过程并改善MCMC收敛后再写入最终排序。",
                        True,
                    ),
                    (
                        "结合前面章节的模拟结果，本文为临床生存预测任务中的模型选择提供了可对照的证据。",
                        False,
                    ),
                ],
            )

    leftovers = [
        "0.8472",
        "0.8429",
        "1.0000",
        "0.1275",
        "0.0893",
        "HR 0.63",
        "HR 1.23",
        "13个Rhat",
        "最小ESS约为10.7",
        "最小ESS约10.7",
        "RSF与RSFLC人数未在性能表中写出",
    ]
    blob = "\n".join(p.text for p in doc.paragraphs)
    for k in leftovers:
        n = blob.count(k)
        if n:
            print("LEFTOVER", k, n)

    doc.save(str(OUT))
    print("saved", OUT)

    doc2 = Document(str(OUT))
    for tbl in doc2.tables:
        h0 = " | ".join(c.text.strip()[:12] for c in tbl.rows[0].cells)
        if "28天AUC" in h0:
            print("T54", h0)
            for row in tbl.rows:
                print(" ", [c.text.strip()[:32] for c in row.cells])
        if tbl.rows[0].cells[0].text.strip() == "纵向标记":
            print("T56")
            for row in tbl.rows:
                print(" ", [c.text.strip()[:36] for c in row.cells])

    red_n = 0
    for p in doc2.paragraphs:
        for r in p.runs:
            if is_red_run(r):
                red_n += 1
    for tbl in doc2.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if is_red_run(r):
                            red_n += 1
    print("red runs", red_n)


if __name__ == "__main__":
    main()
