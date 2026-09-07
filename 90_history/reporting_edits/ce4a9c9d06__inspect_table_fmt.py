# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

p = r"F:\文章_大论文\0722\大论文版本\0826_1_基于机器学习的生存分析方法探讨_第五章RSF_RSFLC图更新.docx"
doc = Document(p)


def borders_info(ti):
    tbl = doc.tables[ti]
    tblPr = tbl._tbl.tblPr
    print("TABLE", ti, tbl.style.name)
    borders = tblPr.find(qn("w:tblBorders")) if tblPr is not None else None
    if borders is not None:
        print(etree.tostring(borders, pretty_print=True, encoding="unicode"))
    else:
        print(" no explicit tblBorders")
    for tag in ["tblW", "jc", "tblLook", "tblCellMar", "tblLayout"]:
        el = tblPr.find(qn(f"w:{tag}")) if tblPr is not None else None
        if el is not None:
            print(tag, etree.tostring(el, encoding="unicode"))
    cell = tbl.rows[0].cells[0]
    tcBorders = cell._tc.tcPr.find(qn("w:tcBorders")) if cell._tc.tcPr is not None else None
    print(" header cell borders:", etree.tostring(tcBorders, encoding="unicode") if tcBorders is not None else None)
    shd = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
    print(" header shading:", etree.tostring(shd, encoding="unicode") if shd is not None else None)
    for ri, ci in [(0, 0), (0, 1), (1, 0), (1, 1)]:
        cell = tbl.rows[ri].cells[ci]
        p0 = cell.paragraphs[0]
        jc = None
        if p0._p.pPr is not None:
            jel = p0._p.pPr.find(qn("w:jc"))
            if jel is not None:
                jc = jel.get(qn("w:val"))
        fonts = []
        for r in p0.runs:
            sz = r.font.size.pt if r.font.size else None
            rPr = r._r.rPr
            ea = None
            ascii_f = None
            if rPr is not None:
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None:
                    ea = rf.get(qn("w:eastAsia"))
                    ascii_f = rf.get(qn("w:ascii"))
            fonts.append((r.text[:25], ascii_f, ea, sz, r.bold))
        print(f"  [{ri},{ci}] jc={jc} fonts={fonts}")


for ti in [4, 5, 6, 7, 8, 9, 10]:
    borders_info(ti)
    print()
