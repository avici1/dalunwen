# -*- coding: utf-8 -*-
"""方案C：将表5-1拆为表5-1a（连续）与表5-1b（分类），并更新正文引用。"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"F:\文章_大论文")
DOCX_IN = ROOT / r"0722\大论文版本\0827_1_基于机器学习的生存分析方法探讨_第五章表格式与表5-1更新.docx"
DOCX_OUT = ROOT / r"0722\大论文版本\0827_1_基于机器学习的生存分析方法探讨_第五章表5-1拆分连续分类.docx"

HEADER_FILL = "D9E2F3"
FONT_ASCII = "Times New Roman"
FONT_EA = "宋体"
FONT_SIZE_PT = 10.5


def set_run_font(run, bold: bool | None = None):
    run.font.name = FONT_ASCII
    run.font.size = Pt(FONT_SIZE_PT)
    if bold is not None:
        run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), FONT_ASCII)
    rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    rFonts.set(qn("w:eastAsia"), FONT_EA)
    rFonts.set(qn("w:cs"), FONT_ASCII)


def set_cell_text(cell, text: str, *, bold: bool = False, align: str = "center"):
    for p in cell.paragraphs:
        for r in p.runs:
            r.text = ""
    p0 = cell.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text
        set_run_font(p0.runs[0], bold=bold)
        for r in p0.runs[1:]:
            r.text = ""
    else:
        run = p0.add_run(text)
        set_run_font(run, bold=bold)

    pPr = p0._p.get_or_add_pPr()
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), align)

    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), "center")


def set_cell_shading(cell, fill: str | None):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tcPr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_tbl_border_el(borders, tag: str):
    el = borders.find(qn(f"w:{tag}"))
    if el is None:
        el = OxmlElement(f"w:{tag}")
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:color"), "000000")
    el.set(qn("w:sz"), "4")
    el.set(qn("w:space"), "0")


def apply_ref_table_format(table, section_names=None):
    section_names = set(section_names or [])
    try:
        table.style = "Normal Table"
    except Exception:
        pass

    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        set_tbl_border_el(borders, edge)

    def upsert(tag, attrs):
        el = tblPr.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tblPr.append(el)
        for k, v in attrs.items():
            el.set(qn(f"w:{k}"), v)

    upsert("tblW", {"w": "5000", "type": "pct"})
    upsert("jc", {"val": "center"})
    upsert("tblLayout", {"type": "autofit"})

    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        tblPr.append(mar)
    for side, w in [("top", "0"), ("left", "108"), ("bottom", "0"), ("right", "108")]:
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), w)
        el.set(qn("w:type"), "dxa")

    n_cols = len(table.columns)
    for ri, row in enumerate(table.rows):
        is_header = ri == 0
        for ci in range(n_cols):
            cell = row.cells[ci]
            text = cell.text.strip()
            if is_header:
                align, bold = "center", True
                set_cell_shading(cell, HEADER_FILL)
            else:
                others = [row.cells[j].text.strip() for j in range(1, n_cols)]
                is_section = text in section_names or (
                    text != "" and all(o in ("", "—", "-") for o in others)
                )
                align = "left" if ci == 0 else "center"
                bold = bool(is_section and ci == 0)
                set_cell_shading(cell, None)
            set_cell_text(cell, text, bold=bold, align=align)


def replace_paragraph_text(paragraph, new_text: str):
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def fill_table(table, header, body_rows):
    """清空并重写已有三列表（行列数不足时扩展）。"""
    need_rows = 1 + len(body_rows)
    while len(table.rows) < need_rows:
        table.add_row()
    while len(table.rows) > need_rows:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)

    for j, h in enumerate(header):
        set_cell_text(table.rows[0].cells[j], h, bold=True, align="center")
    for i, vals in enumerate(body_rows):
        for j, val in enumerate(vals):
            set_cell_text(
                table.rows[i + 1].cells[j],
                val,
                bold=False,
                align="left" if j == 0 else "center",
            )


def insert_table_after_element(doc: Document, anchor_el, header, body_rows, section_names=None):
    table = doc.add_table(rows=1 + len(body_rows), cols=len(header))
    try:
        table.style = "Normal Table"
    except Exception:
        pass
    fill_table(table, header, body_rows)
    anchor_el.addnext(table._tbl)
    apply_ref_table_format(table, section_names=section_names)
    return table


def insert_caption_after_element(style_src_p, anchor_el, text: str):
    new_p_xml = copy.deepcopy(style_src_p._p)
    texts = list(new_p_xml.iter(qn("w:t")))
    if texts:
        texts[0].text = text
        for t in texts[1:]:
            t.text = ""
    else:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        new_p_xml.append(r)
    anchor_el.addnext(new_p_xml)
    return new_p_xml


def find_table_51(doc: Document):
    for ti, tbl in enumerate(doc.tables):
        flat = " ".join(c.text for r in tbl.rows for c in r.cells)
        head = "|".join(c.text.strip() for c in tbl.rows[0].cells)
        if "年龄" in flat and "Charlson" in flat and "样本量" in head and "有观测" not in head:
            return ti, tbl
    raise RuntimeError("未找到表5-1")


def build_split_rows(old_table):
    rows = [[c.text.strip() for c in r.cells] for r in old_table.rows[1:]]
    data = {r[0]: r for r in rows}

    continuous = [
        ["队列规模", "", ""],
        ["总住院次数", "6,299", "—"],
        ["独立患者数", "6,092", "—"],
        ["人口学与入院特征", "", ""],
    ]
    cont_vars = [
        "年龄（岁）",
        "体重指数（BMI）",
        "Charlson合并症指数",
        "ICU住院天数",
        "总住院天数",
        "APSIII",
        "OASIS",
        "SAPSII",
    ]
    for v in cont_vars:
        if v == "APSIII":
            continuous.append(["入院严重度评分", "", ""])
        if v in data:
            continuous.append([data[v][0], data[v][1], data[v][2]])

    categorical = [
        ["人口学与分类特征", "", ""],
    ]
    cat_vars = [
        "男性，例(%)",
        "择期手术，例(%)",
        "机械通气，例(%)",
        "缺血性卒中（I63脑梗死），例(%)",
        "脑出血（I61），例(%)",
        "蛛网膜下腔出血（I60），例(%)",
        "混合型，例(%)",
        "院内死亡，例(%)",
        "ICU内死亡，例(%)",
        "28天死亡，例(%)",
        "90天死亡，例(%)",
    ]
    for v in cat_vars:
        if v == "缺血性卒中（I63脑梗死），例(%)":
            categorical.append(["卒中分型", "", ""])
        if v == "院内死亡，例(%)":
            categorical.append(["结局（供参考）", "", ""])
        if v in data:
            categorical.append([data[v][0], data[v][1], data[v][2]])

    return continuous, categorical


def main():
    shutil.copy2(DOCX_IN, DOCX_OUT)
    doc = Document(str(DOCX_OUT))

    cap_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("表5-1 ") or p.text.strip() == "表5-1 研究队列基线数据基本特征":
            cap_idx = i
            break
    if cap_idx is None:
        raise RuntimeError("未找到表5-1标题")

    _, old_tbl = find_table_51(doc)
    continuous, categorical = build_split_rows(old_tbl)

    for p in doc.paragraphs:
        t = p.text
        if "表5-1汇总基线数据基本特征" in t or (t.startswith("基线特征包含") and "表5-1" in t):
            new = (
                "基线特征包含年龄、性别、体重指数（BMI）、卒中分型、Charlson合并症指数（CCI），"
                "三个常用的重症评分系统即急性生理学评分III（APSIII）、牛津急性严重度评分（OASIS）、"
                "简化急性生理学评分II（SAPSII），以及机械通气、择期手术和住院时长等指标。"
                "纵向变量包括每日测量的格拉斯哥昏迷评分（GCS）、序贯器官衰竭评估（SOFA）评分、血肌酐、乳酸、"
                "动脉血气指标（pH、PaO₂、PaCO₂、PaO₂/FiO₂比值），以及血红蛋白、血糖、血钠、碳酸氢盐等实验室与生命体征指标。"
                "纵向数据约78,649条患者-日记录，覆盖6,299次住院（6,092名患者）。"
                "为便于阅读，基线特征按变量类型分列两表："
                "表5-1a汇总连续型变量（均值±标准差），表5-1b汇总分类型变量（例数及构成比）；"
                "表5-2按方案先在住院水平计算第1天值、住院期均值与最差值，再在队列水平报告均值±标准差，"
                "以避免重复测量导致的样本量膨胀。主要结局分布列入表5-1b供参考。"
            )
            replace_paragraph_text(p, new)
            print("已更新正文引用段")
            break

    cap = doc.paragraphs[cap_idx]
    replace_paragraph_text(cap, "表5-1a 研究队列基线连续变量特征")
    fill_table(old_tbl, ["变量", "样本量", "均值±标准差"], continuous)
    apply_ref_table_format(
        old_tbl,
        section_names={"队列规模", "人口学与入院特征", "入院严重度评分"},
    )

    style_src = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("表5-2"):
            style_src = p
            break
    if style_src is None:
        style_src = cap

    # 顺序：表5-1a → 表5-1b标题 → 表5-1b → 表5-2...
    # 先在表后插入表格B，再在表A与表B之间插入标题B
    tbl_b = insert_table_after_element(
        doc,
        old_tbl._tbl,
        ["变量", "样本量", "例数(%)"],
        categorical,
        section_names={"人口学与分类特征", "卒中分型", "结局（供参考）"},
    )
    insert_caption_after_element(
        style_src, old_tbl._tbl, "表5-1b 研究队列基线分类变量特征"
    )
    # 上面两步后顺序应为：表A → 标题B → 表B（因第二次 addnext 插在表A后、表B前）

    for p in doc.paragraphs:
        t = p.text
        if "表5-1" in t and "表5-1a" not in t and "表5-1b" not in t:
            nt = t.replace("表5-1汇总基线数据基本特征", "表5-1a与表5-1b汇总基线数据基本特征")
            nt = nt.replace("列入表5-1供参考", "列入表5-1b供参考")
            nt = nt.replace("见表5-1", "见表5-1a与表5-1b")
            nt = nt.replace("如表5-1", "如表5-1a与表5-1b")
            if nt != t:
                replace_paragraph_text(p, nt)
                print("额外替换:", t[:60], "->", nt[:60])

    doc.save(str(DOCX_OUT))
    print("saved:", DOCX_OUT)
    _ = tbl_b

    doc2 = Document(str(DOCX_OUT))
    print("==== captions / cites ====")
    for i, p in enumerate(doc2.paragraphs):
        t = p.text.strip()
        if "表5-1" in t or t.startswith("表5-2"):
            print(i, t[:200])

    # 核对正文顺序：caption a / table a / caption b / table b / caption 5-2
    from docx.oxml.ns import qn as _qn

    body = doc2.element.body
    seq = []
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            texts = [t.text or "" for t in child.iter(_qn("w:t"))]
            txt = "".join(texts).strip()
            if txt.startswith("表5-1") or txt.startswith("表5-2"):
                seq.append(("p", txt[:60]))
        elif tag == "tbl":
            first = ""
            for t in child.iter(_qn("w:t")):
                if t.text and t.text.strip():
                    first = t.text.strip()
                    break
            if first in ("变量",):
                # peek second cell of header via collecting first row texts briefly
                row0 = []
                for tr in child.iter(_qn("w:tr")):
                    for tc in tr.iter(_qn("w:tc")):
                        row0.append("".join((x.text or "") for x in tc.iter(_qn("w:t"))).strip())
                    break
                seq.append(("tbl", "|".join(row0[:3])))
                if len(seq) >= 6:
                    break
    print("==== body order around ch5 tables ====")
    for item in seq[:8]:
        print(item)

    print("==== 5-1a ====")
    for tbl in doc2.tables:
        head = "|".join(c.text.strip() for c in tbl.rows[0].cells)
        if head == "变量|样本量|均值±标准差":
            for r in tbl.rows:
                print(" | ".join(c.text.strip() for c in r.cells))
        if head == "变量|样本量|例数(%)":
            print("==== 5-1b ====")
            for r in tbl.rows:
                print(" | ".join(c.text.strip() for c in r.cells))


if __name__ == "__main__":
    main()
