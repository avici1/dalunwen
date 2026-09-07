# -*- coding: utf-8 -*-
"""合并 bootstrap_summary2（或旧版 boostrap_summary）与 friedman_pairwise，写入带分组表头的 xlsx。"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

base = Path(__file__).resolve().parent
out_path = base / "result_Friedman_wilcoxon_boostrap.xlsx"

_sum = base / "bootstrap_summary2.xlsx"
_bleg = base / "boostrap_summary.xlsx"
b = pd.read_excel(_sum if _sum.exists() else _bleg, sheet_name="summary")
f = pd.read_excel(base / "friedman_pairwise.xlsx")
merged = b.merge(f, on=["metric", "contrast"], how="left", validate="one_to_one")

cols = [
    "metric",
    "contrast",
    "V",
    "p.value",
    "p.holm",
    "model_left",
    "model_right",
    "n_folds",
    "mean_left",
    "mean_right",
    "mean_diff_left_minus_right",
    "sd_diff",
    "CI_percentile_2.5",
    "CI_percentile_97.5",
    "CI_BCa_2.5",
    "CI_BCa_97.5",
    "p_wilcoxon_paired",
    "p_bootstrap",
]
merged = merged[cols]

wb = Workbook()
ws = wb.active
ws.title = "合并结果"

# 第 1 行：分组说明（横向合并）
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
ws.cell(1, 1).value = (
    "（一）Friedman 检验：四模型整体差异的 Friedman 检验见单独表「friedman_omnibus」；"
    "本表两两比较与之一致（折=区组，指标均为各折测试集表现）。"
)
ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=5)
ws.cell(1, 3).value = (
    "（二）Wilcoxon 检验：Friedman 事后、折内配对的 Wilcoxon 符号秩检验；"
    "p.value 为原始双侧 p，p.holm 为对全部两两比较作 Holm 多重比较校正后的 p。"
)
ws.merge_cells(start_row=1, start_column=6, end_row=1, end_column=18)
ws.cell(1, 6).value = (
    "（三）Bootstrap 检验：对两模型逐折测试集指标之差（左-右）做有放回重抽样（R=10000），"
    "报告均值差、折间差之 SD、百分位法与 BCa 的 95%CI、bootstrap p（p_bootstrap）；"
    "并附未作 Holm 校正的配对 Wilcoxon p（与左侧 Wilcoxon 校正策略不同，请勿混读为同一检验）。"
)

hdr_align = Alignment(
    horizontal="center", vertical="center", wrap_text=True
)
title_font = Font(size=10)
for c in (1, 3, 6):
    cell = ws.cell(1, c)
    cell.alignment = hdr_align
    cell.font = title_font

ws.row_dimensions[1].height = 72

# 第 2 行：列名
for j, name in enumerate(cols, start=1):
    c = ws.cell(2, j)
    c.value = name
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 数据从第 3 行
for i, row in enumerate(merged.itertuples(index=False), start=3):
    for j, val in enumerate(row, start=1):
        ws.cell(i, j).value = val

# 列宽
for j in range(1, len(cols) + 1):
    ws.column_dimensions[get_column_letter(j)].width = min(18, 12 + len(cols[j - 1]) * 0.3)

ws.freeze_panes = "A3"
wb.save(out_path)
print("已保存:", out_path)
